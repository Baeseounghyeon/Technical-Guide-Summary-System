"""견적서 만들기.

- 양식(Word·Excel)을 고르면 양식 안의 자리 표시를 채운다.
    글자: {{제목}} {{작성자}} {{소속}} {{부서}} {{이메일}} {{전화번호}} {{작성일}} {{버전}}
    표:   {{제품가격표}} {{매입금액}}  — 그 자리(Word 는 그 문단, Excel 은 그 칸부터 아래로)에 체크한 항목 표를 넣는다.
    표 자리 표시가 없으면 문서 끝(Excel 은 첫 시트의 마지막 행 아래)에 붙인다.
- 양식이 없으면 기본 양식(규격서와 같은 남색 띠 형식)으로 PDF·Word·Excel 을 만든다.
- 미리보기는 실제로 만들어질 파일을 HTML 로 바꿔 보여 준다.
"""

import html
import io
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .docx_export import _set_default_font
from .form_export import NAVY, _shade

AUTHOR_FIELDS = [("name", "작성자"), ("company", "소속"), ("department", "부서"), ("email", "이메일"), ("phone", "전화번호")]
MEDIA = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


@dataclass
class ItemTable:
    label: str  # 제품 가격표 / 매입 금액
    columns: list[str]
    rows: list[list[str]]


@dataclass
class QuoteData:
    title: str
    author: dict[str, str]
    tables: list[ItemTable]
    date: str
    version: str = ""
    meta: list[tuple[str, str]] = field(default_factory=list)

    def text_values(self) -> dict[str, str]:
        values = {"{{제목}}": self.title, "{{작성일}}": self.date, "{{버전}}": self.version}
        for key, label in AUTHOR_FIELDS:
            values["{{" + label + "}}"] = self.author.get(key, "")
        return values

    def table_by_key(self) -> dict[str, ItemTable]:
        keys = {"제품 가격표": "{{제품가격표}}", "매입 금액": "{{매입금액}}"}
        return {keys[t.label]: t for t in self.tables if t.label in keys}


def formats_for(template_ext: str | None) -> list[str]:
    """받을 수 있는 형식: 양식 형식이 먼저, PDF 는 기본 양식으로 만든다."""
    if template_ext == "docx":
        return ["docx", "pdf"]
    if template_ext == "xlsx":
        return ["xlsx", "pdf"]
    return ["pdf", "docx", "xlsx"]


# ---------- 기본 양식 ----------


def _html_table(t: ItemTable) -> str:
    head = "".join(f'<th style="border:1px solid #9AA3B5;background-color:#D9DFEA;padding:3px">{html.escape(c)}</th>' for c in t.columns)
    body = "".join(
        "<tr>" + "".join(f'<td style="border:1px solid #9AA3B5;padding:3px">{html.escape(c)}</td>' for c in r) + "</tr>" for r in t.rows
    )
    return f'<table style="width:100%;border-collapse:collapse;font-size:8.5pt"><tr>{head}</tr>{body}</table>'


def builtin_html(q: QuoteData) -> str:
    author = "".join(
        f'<tr><td style="width:22%;font-weight:bold;padding:2px 6px">{label}</td><td style="padding:2px 6px">{html.escape(q.author.get(k, "") or "-")}</td></tr>'
        for k, label in AUTHOR_FIELDS
    )
    tables = "".join(
        f'<p style="font-weight:bold;font-size:10pt;margin:12px 0 4px 0">{html.escape(t.label)}</p>{_html_table(t)}' for t in q.tables
    )
    meta = " · ".join(f"{html.escape(k)}: {html.escape(v)}" for k, v in q.meta)
    return (
        f'<p style="text-align:center;font-size:16pt;margin:0 0 4px 0">{html.escape(q.title)}</p>'
        f'<p style="text-align:center;font-size:9pt;color:#555;margin:0 0 12px 0">견적서 · {html.escape(q.date)}</p>'
        f'<div style="background-color:#E6EBF2;padding:10px">'
        f'<p style="background-color:#{NAVY};color:#FFFFFF;text-align:center;font-size:12pt;font-weight:bold;padding:6px;margin:0 0 8px 0">견적서 정보</p>'
        f'<table style="width:100%;font-size:9.5pt">{author}</table>{tables}</div>'
        f'<p style="font-size:8pt;color:#666;margin-top:10px">{meta}</p>'
    )


def builtin_pdf(q: QuoteData) -> bytes:
    story = pymupdf.Story(html=builtin_html(q), user_css="body{font-size:10pt;}")
    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    page = pymupdf.paper_rect("a4")
    area = page + (40, 40, -40, -40)
    more = True
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(area)
        story.draw(device)
        writer.end_page()
    writer.close()
    with pymupdf.open("pdf", buffer.getvalue()) as pdf:
        pdf.subset_fonts()
        return pdf.tobytes(garbage=4, deflate=True)


def _docx_table(doc, t: ItemTable) -> Table:
    table = doc.add_table(rows=1, cols=max(len(t.columns), 1))
    table.style = "Table Grid"
    for cell, name in zip(table.rows[0].cells, t.columns):
        _shade(cell, "D9DFEA")
        run = cell.paragraphs[0].add_run(name)
        run.bold = True
        run.font.size = Pt(9)
    for r in t.rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, r):
            cell.paragraphs[0].add_run(value).font.size = Pt(9)
    return table


def builtin_docx(q: QuoteData) -> bytes:
    doc = Document()
    _set_default_font(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(q.title)
    run.bold = True
    run.font.size = Pt(16)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"견적서 · {q.date}").font.size = Pt(9)
    bar = doc.add_table(rows=1, cols=1).rows[0].cells[0]
    _shade(bar, NAVY)
    r = bar.paragraphs[0].add_run("견적서 정보")
    r.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    bar.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    info = doc.add_table(rows=0, cols=2)
    info.style = "Table Grid"
    for k, label in AUTHOR_FIELDS:
        cells = info.add_row().cells
        cells[0].paragraphs[0].add_run(label).bold = True
        cells[1].paragraphs[0].add_run(q.author.get(k, "") or "-")
    for t in q.tables:
        doc.add_paragraph().add_run(t.label).bold = True
        _docx_table(doc, t)
    if q.meta:
        doc.add_paragraph().add_run(" · ".join(f"{k}: {v}" for k, v in q.meta)).font.size = Pt(8)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _xlsx_write_table(ws, row: int, col: int, t: ItemTable) -> int:
    """(row, col)부터 머리글+행을 쓰고 다음 빈 행 번호를 돌려준다."""
    fill = PatternFill("solid", fgColor="D9DFEA")
    for j, name in enumerate(t.columns):
        cell = ws.cell(row=row, column=col + j, value=name)
        cell.font = Font(bold=True)
        cell.fill = fill
    for i, r in enumerate(t.rows, start=1):
        for j, value in enumerate(r):
            ws.cell(row=row + i, column=col + j, value=_number(value))
    return row + len(t.rows) + 1


def _number(value: str):
    """'1,200,000' 같은 숫자는 숫자 칸으로 넣어 Excel 에서 계산할 수 있게 한다."""
    plain = value.replace(",", "").strip()
    if not any(ch.isdigit() for ch in plain):
        return value
    for convert in (int, float):
        try:
            return convert(plain)
        except ValueError:
            continue
    return value


def builtin_xlsx(q: QuoteData) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "견적서"
    ws.cell(row=1, column=1, value=q.title).font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=f"견적서 · {q.date}")
    row = 4
    for k, label in AUTHOR_FIELDS:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=q.author.get(k, ""))
        row += 1
    for t in q.tables:
        row += 1
        ws.cell(row=row, column=1, value=t.label).font = Font(bold=True, color=NAVY)
        row = _xlsx_write_table(ws, row + 1, 1, t)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------- 양식 채우기 ----------


def _replace_in_paragraph(p: Paragraph, values: dict[str, str]) -> None:
    text = p.text
    if "{{" not in text:
        return
    new = text
    for key, value in values.items():
        new = new.replace(key, value)
    if new != text and p.runs:
        p.runs[0].text = new
        for run in p.runs[1:]:
            run.text = ""


def _all_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from cell.paragraphs


def fill_docx(template: Path, q: QuoteData) -> bytes:
    doc = Document(str(template))
    tables = q.table_by_key()
    placed: set[str] = set()
    # 표 자리 표시 문단 → 그 자리에 표
    for p in list(doc.paragraphs):
        for key, t in tables.items():
            if key in p.text:
                table = _docx_table(doc, t)
                p._p.addnext(table._tbl)
                p.runs[0].text = p.text.replace(key, "").strip() if p.runs else ""
                for run in p.runs[1:]:
                    run.text = ""
                placed.add(key)
    values = q.text_values()
    for p in _all_paragraphs(doc):
        _replace_in_paragraph(p, values)
    for key, t in tables.items():
        if key not in placed:
            doc.add_paragraph().add_run(t.label).bold = True
            _docx_table(doc, t)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def fill_xlsx(template: Path, q: QuoteData) -> bytes:
    wb = load_workbook(str(template))
    values = q.text_values()
    tables = q.table_by_key()
    anchors: dict[str, tuple] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str) or "{{" not in cell.value:
                    continue
                stripped = cell.value.strip()
                if stripped in tables:
                    anchors[stripped] = (ws, cell.row, cell.column)
                    cell.value = None
                    continue
                new = cell.value
                for key, value in values.items():
                    new = new.replace(key, value)
                cell.value = new
    for key, t in tables.items():
        if key in anchors:
            ws, r, c = anchors[key]
            _xlsx_write_table(ws, r, c, t)
        else:
            ws = wb.worksheets[0]
            r = ws.max_row + 2
            ws.cell(row=r, column=1, value=t.label).font = Font(bold=True)
            _xlsx_write_table(ws, r + 1, 1, t)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------- 미리보기 HTML ----------

_PREVIEW_CSS = (
    "<style>body{font-family:'Malgun Gothic',sans-serif;font-size:13px;color:#1d2433;margin:16px}"
    "table{border-collapse:collapse;margin:6px 0}td,th{border:1px solid #c5ccd8;padding:3px 6px;vertical-align:top}"
    "h3{margin:16px 0 6px}</style>"
)


def docx_to_html(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = []
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            text = Paragraph(child, doc).text
            parts.append(f"<p>{html.escape(text) or '&nbsp;'}</p>")
        elif tag == "tbl":
            t = Table(child, doc)
            rows = "".join("<tr>" + "".join(f"<td>{html.escape(c.text)}</td>" for c in r.cells) + "</tr>" for r in t.rows)
            parts.append(f"<table>{rows}</table>")
    return _PREVIEW_CSS + "".join(parts)


def xlsx_to_html(data: bytes) -> str:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 500), max_col=min(ws.max_column, 30)):
            cells = "".join(f"<td>{html.escape('' if c.value is None else str(c.value))}</td>" for c in row)
            rows.append(f"<tr>{cells}</tr>")
        parts.append(f"<h3>{html.escape(ws.title)}</h3><table>{''.join(rows)}</table>")
    return _PREVIEW_CSS + "".join(parts)


def build_quote(q: QuoteData, fmt: str, template: Path | None, template_ext: str | None) -> bytes:
    """fmt 형식으로 만든다. 양식과 같은 형식이면 양식을 채우고, 아니면 기본 양식."""
    if template is not None and fmt == template_ext == "docx":
        return fill_docx(template, q)
    if template is not None and fmt == template_ext == "xlsx":
        return fill_xlsx(template, q)
    return {"pdf": builtin_pdf, "docx": builtin_docx, "xlsx": builtin_xlsx}[fmt](q)


def preview_html(q: QuoteData, template: Path | None, template_ext: str | None) -> str:
    """대표 형식(양식이 있으면 양식 형식)으로 만든 결과를 HTML 로."""
    if template is not None and template_ext == "docx":
        return docx_to_html(fill_docx(template, q))
    if template is not None and template_ext == "xlsx":
        return xlsx_to_html(fill_xlsx(template, q))
    return _PREVIEW_CSS + builtin_html(q)


__all__ = ["AUTHOR_FIELDS", "MEDIA", "ItemTable", "QuoteData", "build_quote", "formats_for", "preview_html"]
