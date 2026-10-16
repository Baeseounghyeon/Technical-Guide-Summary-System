"""규격서·견적서 생성 결과를 PDF·Word 로 만든다.

양식: 가운데 제목 → 남색 띠("○○ 정보") → 두 칸 배치의 항목(이름 위, 값 상자 아래).
긴 항목(wide)은 두 칸을 모두 쓴다. 아래에 작성자·소속·작성일·버전을 적는다.
"""

import html
import io
from dataclasses import dataclass

import pymupdf
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .docx_export import _set_default_font

NAVY = "1F2A6B"


@dataclass
class FormDoc:
    kind: str  # 규격서 / 견적서
    title: str
    fields: list[dict]  # FORM_FIELDS 항목
    values: dict[str, str]
    meta: list[tuple[str, str]]  # (작성자, …)


def _rows(form: FormDoc) -> list[list[dict]]:
    """두 칸씩 묶는다. wide 항목은 한 줄을 혼자 쓴다."""
    rows: list[list[dict]] = []
    pending: list[dict] = []
    for f in form.fields:
        if f.get("wide"):
            if pending:
                rows.append(pending)
                pending = []
            rows.append([f])
        else:
            pending.append(f)
            if len(pending) == 2:
                rows.append(pending)
                pending = []
    if pending:
        rows.append(pending)
    return rows


# ---------- PDF ----------


def build_pdf(form: FormDoc) -> bytes:
    def cell(f: dict, span: int) -> str:
        value = html.escape(form.values.get(f["key"], "")).replace("\n", "<br/>")
        return (
            f'<td colspan="{span}" style="width:{50 * span}%;padding:4px 8px;vertical-align:top">'
            f'<p style="font-size:8pt;font-weight:bold;margin:0 0 2px 0">{html.escape(f["label"])}:</p>'
            f'<p style="border:1px solid #9AA3B5;background-color:#FFFFFF;padding:5px;margin:0;font-size:9.5pt">{value or "&#160;"}</p>'
            "</td>"
        )

    rows = "".join(
        "<tr>" + "".join(cell(f, 2 if len(r) == 1 and f.get("wide") else 1) for f in r) + ("<td></td>" if len(r) == 1 and not r[0].get("wide") else "") + "</tr>"
        for r in _rows(form)
    )
    meta = " · ".join(f"{html.escape(k)}: {html.escape(v)}" for k, v in form.meta)
    body = (
        f'<p style="text-align:center;font-size:15pt;margin:0 0 4px 0">{html.escape(form.title)}</p>'
        f'<p style="text-align:center;font-size:9pt;color:#555;margin:0 0 14px 0">{html.escape(form.kind)}</p>'
        f'<div style="background-color:#E6EBF2;padding:10px">'
        f'<p style="background-color:#{NAVY};color:#FFFFFF;text-align:center;font-size:12pt;font-weight:bold;padding:6px;margin:0 0 8px 0">'
        f"{html.escape(form.kind)} 정보</p>"
        f'<table style="width:100%">{rows}</table>'
        "</div>"
        f'<p style="font-size:8pt;color:#666;margin-top:10px">{meta}</p>'
    )
    story = pymupdf.Story(html=body, user_css="body{font-size:10pt;}")
    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    page = pymupdf.paper_rect("a4")
    area = page + (50, 50, -50, -50)
    more = True
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(area)
        story.draw(device)
        writer.end_page()
    writer.close()
    with pymupdf.open("pdf", buffer.getvalue()) as pdf:
        pdf.subset_fonts()  # 한글 글꼴은 쓰인 글자만 남긴다
        return pdf.tobytes(garbage=4, deflate=True)


# ---------- Word ----------


def _shade(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    props.append(shd)


def build_docx(form: FormDoc) -> bytes:
    doc = Document()
    _set_default_font(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(form.title)
    run.bold = True
    run.font.size = Pt(16)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(form.kind).font.size = Pt(9)

    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    head = table.add_row().cells
    bar = head[0].merge(head[1])
    _shade(bar, NAVY)
    p = bar.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{form.kind} 정보")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for row_fields in _rows(form):
        cells = table.add_row().cells
        targets = [cells[0].merge(cells[1])] if len(row_fields) == 1 and row_fields[0].get("wide") else cells
        for target, f in zip(targets, row_fields):
            label = target.paragraphs[0].add_run(f"{f['label']}:")
            label.bold = True
            label.font.size = Pt(8)
            for line in (form.values.get(f["key"], "") or "").splitlines() or [""]:
                target.add_paragraph().add_run(line).font.size = Pt(10)
    meta = doc.add_paragraph()
    m = meta.add_run(" · ".join(f"{k}: {v}" for k, v in form.meta))
    m.font.size = Pt(8)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


MEDIA = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def build_form_file(form: FormDoc, fmt: str) -> bytes:
    return build_pdf(form) if fmt == "pdf" else build_docx(form)
