"""등록한 원본 문서를 다른 형식(Word·PowerPoint·Excel·PDF)으로 내려받기.

분석한 목차별 본문(sections)으로 새 파일을 만든다. 원본의 그림·표 서식은 담기지 않는다
(원본 그대로는 '원본 파일' 받기를 쓴다).
"""

import html
import io
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import pymupdf
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches
from pptx.util import Pt as PptPt

from .docx_export import _heading, _set_default_font

MEDIA_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}
SLIDE_CHARS = 900  # 슬라이드 한 장에 넣을 본문 글자 수
EXCEL_CELL_LIMIT = 32000  # 엑셀 셀 최대 32,767자


@dataclass
class DocSection:
    title: str
    level: int
    text: str
    page_start: int | None
    page_end: int | None


@dataclass
class DocInfo:
    title: str
    meta: list[tuple[str, str]]  # (항목, 값) — 벤더, 종류, 설명, 등록자, 등록일
    sections: list[DocSection]


def _pages(s: DocSection) -> str:
    if not s.page_start:
        return ""
    return f"p.{s.page_start}" + (f"–{s.page_end}" if s.page_end and s.page_end != s.page_start else "")


def build_document_file(info: DocInfo, fmt: str, out_path: Path) -> None:
    {"docx": _docx, "pptx": _pptx, "xlsx": _xlsx, "pdf": _pdf}[fmt](info, out_path)


def _docx(info: DocInfo, out_path: Path) -> None:
    doc = Document()
    _set_default_font(doc)
    doc.add_paragraph(info.title, style="Title")
    for key, value in info.meta:
        p = doc.add_paragraph()
        p.add_run(f"{key}: ").bold = True
        p.add_run(value)
    for s in info.sections:
        _heading(doc, s.title, min(max(s.level, 1), 4))
        for line in s.text.splitlines():
            if line.strip():
                doc.add_paragraph(line.strip())
        if label := _pages(s):
            run = doc.add_paragraph().add_run(f"원본 {label}")
            run.italic = True
            run.font.size = Pt(8)
    doc.save(str(out_path))


def _pptx(info: DocInfo, out_path: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = info.title
    cover.placeholders[1].text = "\n".join(f"{k}: {v}" for k, v in info.meta)
    layout = prs.slide_layouts[1]  # 제목 + 본문
    for s in info.sections:
        text = " ".join(s.text.split())
        chunks = [text[i : i + SLIDE_CHARS] for i in range(0, len(text), SLIDE_CHARS)] or [""]
        for index, chunk in enumerate(chunks):
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = s.title + (" (계속)" if index else "")
            body = slide.placeholders[1].text_frame
            body.text = chunk + (f"\n({_pages(s)})" if _pages(s) and index == len(chunks) - 1 else "")
            body.word_wrap = True
            body.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            for paragraph in body.paragraphs:
                for run in paragraph.runs:
                    run.font.size = PptPt(14)
    prs.save(str(out_path))


def _xlsx(info: DocInfo, out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "목차별 본문"
    ws.append(["No", "단계", "목차", "원본 쪽", "글자 수", "본문"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, s in enumerate(info.sections, start=1):
        ws.append([i, s.level, s.title, _pages(s), len(s.text), s.text[:EXCEL_CELL_LIMIT]])
    for col, width in zip("ABCDEF", (6, 6, 40, 12, 9, 100)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    info_ws = wb.create_sheet("문서 정보")
    info_ws.append(["제목", info.title])
    for key, value in info.meta:
        info_ws.append([key, value])
    info_ws.column_dimensions["A"].width = 12
    info_ws.column_dimensions["B"].width = 80
    wb.save(str(out_path))


def _pdf(info: DocInfo, out_path: Path) -> None:
    parts = [f"<h1>{html.escape(info.title)}</h1>"]
    parts += [f"<p><b>{html.escape(k)}:</b> {html.escape(v)}</p>" for k, v in info.meta]
    for s in info.sections:
        level = min(max(s.level, 1), 4) + 1
        parts.append(f"<h{level}>{html.escape(s.title)}</h{level}>")
        parts += [f"<p>{html.escape(line.strip())}</p>" for line in s.text.splitlines() if line.strip()]
        if label := _pages(s):
            parts.append(f'<p style="font-size:8pt;color:#666"><i>원본 {label}</i></p>')
    story = pymupdf.Story(html="\n".join(parts), user_css="body{font-size:10pt;} h1{font-size:18pt;}")
    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    page = pymupdf.paper_rect("a4")
    area = page + (50, 50, -50, -50)
    more = True
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(area)
        story.draw(device)
        writer.end_page()
    writer.close()
    # 한글 글꼴 전체가 들어가 파일이 커지므로 쓰인 글자만 남긴다.
    with pymupdf.open("pdf", buffer.getvalue()) as pdf:
        pdf.subset_fonts()
        pdf.save(str(out_path), garbage=4, deflate=True)


def doc_info(doc, doc_type_label: str) -> DocInfo:
    """SourceDocument → 내려받기용 정보."""
    created: datetime = doc.created_at
    meta = [("벤더", doc.vendor.name), ("문서 종류", doc_type_label)]
    if doc.description:
        meta.append(("설명", doc.description))
    meta += [("등록자", doc.uploader.name if doc.uploader else "-"), ("등록일", f"{created:%Y-%m-%d}")]
    sections = [DocSection(s.title, s.level, s.text or "", s.page_start, s.page_end) for s in doc.sections]
    return DocInfo(title=doc.title or doc.url or f"문서 {doc.id}", meta=meta, sections=sections)
