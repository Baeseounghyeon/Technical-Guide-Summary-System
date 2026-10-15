"""파트너 가이드 수정 결과 파일: 변경 내용 추적 Word, 수정 보고서 Word."""

import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

from .docx_export import _set_default_font

CATEGORY_LABELS = {
    "wrong_value": "잘못된 값",
    "wrong_command": "잘못된 명령",
    "missing_step": "누락된 절차",
    "wrong_order": "순서 오류",
    "missing_warning": "누락된 주의사항",
    "contradiction": "원본과 모순",
    "outdated": "구버전 내용",
}
SEVERITY_LABELS = {"high": "높음", "medium": "보통", "low": "낮음"}
STATUS_LABELS = {"open": "미결정", "accepted": "수락", "rejected": "거절"}

_CONTENT_TAGS = {qn("w:r"), qn("w:hyperlink"), qn("w:ins"), qn("w:del")}


@dataclass
class TrackedChange:
    finding_id: int
    before: str
    after: str


def build_tracked_docx(
    src: Path, changes: list[TrackedChange], author: str, when: datetime, out_path: Path
) -> set[int]:
    """수락된 수정을 Word '변경 내용 추적'으로 반영한다. 반영한 finding id 를 돌려준다.

    수정 전 문구가 한 문단 안에 그대로 있어야 반영된다. 반영한 문단은 서식이 첫 글자 서식으로 통일되고
    문단 안의 하이퍼링크는 일반 글자로 바뀐다.
    """
    doc = Document(str(src))
    pending = list(changes)
    applied: set[int] = set()
    rev_id = 1000

    for paragraph in _all_paragraphs(doc):
        if not pending:
            break
        text = paragraph.text
        edits: list[tuple[int, int, str, int]] = []
        for change in list(pending):
            start = text.find(change.before) if change.before else -1
            if start < 0:
                continue
            end = start + len(change.before)
            if any(start < e and s < end for s, e, _, _ in edits):
                continue  # 같은 문단의 다른 수정과 겹치면 건너뛴다
            edits.append((start, end, change.after, change.finding_id))
            pending.remove(change)
        if edits:
            rev_id = _rewrite_with_revisions(paragraph, text, sorted(edits), author, when, rev_id)
            applied.update(fid for *_, fid in edits)

    doc.save(str(out_path))
    return applied


def _all_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _rewrite_with_revisions(
    paragraph: Paragraph, text: str, edits: list[tuple[int, int, str, int]], author: str, when: datetime, rev_id: int
) -> int:
    p = paragraph._p
    first_run = p.find(qn("w:r"))
    rpr = first_run.find(qn("w:rPr")) if first_run is not None else None
    for child in list(p):
        if child.tag in _CONTENT_TAGS:
            p.remove(child)

    date = when.strftime("%Y-%m-%dT%H:%M:%SZ")
    pos = 0
    for start, end, after, _ in edits:
        if start > pos:
            p.append(_run(text[pos:start], rpr))
        deleted = OxmlElement("w:del")
        _revision_attrs(deleted, rev_id, author, date)
        deleted.append(_run(text[start:end], rpr, deleted=True))
        p.append(deleted)
        inserted = OxmlElement("w:ins")
        _revision_attrs(inserted, rev_id + 1, author, date)
        inserted.append(_run(after, rpr))
        p.append(inserted)
        rev_id += 2
        pos = end
    if pos < len(text):
        p.append(_run(text[pos:], rpr))
    return rev_id


def _revision_attrs(el, rev_id: int, author: str, date: str) -> None:
    el.set(qn("w:id"), str(rev_id))
    el.set(qn("w:author"), author)
    el.set(qn("w:date"), date)


def _run(text: str, rpr, *, deleted: bool = False):
    r = OxmlElement("w:r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    t = OxmlElement("w:delText" if deleted else "w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


@dataclass
class ReportRow:
    section: str
    category: str
    severity: str
    before: str
    after: str
    reason: str
    evidence: str
    status: str
    decided: str  # 결정자 · 일시
    note: str


def build_report_docx(
    *, partner_title: str, vendor: str, guide_title: str, created_at: datetime, rows: list[ReportRow], out_path: Path
) -> None:
    doc = Document()
    _set_default_font(doc)
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width

    doc.add_paragraph("파트너 가이드 수정 보고서", style="Title")
    doc.add_paragraph(f"검토 대상: {partner_title}")
    doc.add_paragraph(f"기준 문서: {vendor} · {guide_title}")
    doc.add_paragraph(f"작성일: {created_at:%Y-%m-%d %H:%M}")
    counts = {s: sum(r.status == s for r in rows) for s in STATUS_LABELS}
    doc.add_paragraph(
        f"오류 후보 {len(rows)}건 · 수락 {counts['accepted']} · 거절 {counts['rejected']} · 미결정 {counts['open']}"
    )

    headers = ["No", "위치", "유형", "심각도", "수정 전", "수정 후", "사유", "벤더 원본 근거", "결정", "결정자·일시", "메모"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = label
    for no, r in enumerate(rows, start=1):
        values = [
            str(no), r.section, CATEGORY_LABELS.get(r.category, r.category), SEVERITY_LABELS.get(r.severity, r.severity),
            r.before, r.after, r.reason, r.evidence, STATUS_LABELS.get(r.status, r.status), r.decided, r.note,
        ]
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)

    doc.save(str(out_path))
