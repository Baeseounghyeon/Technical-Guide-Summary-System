import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

from ..llm.base import language_native_name
from .base import REVISION_COLUMNS, ExportContext, open_template, page_label

KOREAN_FONT = "맑은 고딕"
CODE_FONT = "Consolas"

_BULLET = re.compile(r"^\s*[-*+]\s+(.*)")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)")
_SUBHEADING = re.compile(r"^\s*#{1,6}\s+(.*)")
_INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def build_docx(ctx: ExportContext, out_path: Path) -> None:
    if ctx.template_path:
        # 파트너 양식: 본문 내용은 지우고 디자인(스타일, 머리글·바닥글의 로고, 용지·여백)만 쓴다.
        doc = Document(open_template(ctx.template_path))
        _clear_body(doc)
    else:
        doc = Document()
        _set_default_font(doc)

    # 표지
    if ctx.logo_path:
        doc.add_picture(str(ctx.logo_path), width=Cm(5))
        if not ctx.template_path:  # 양식에는 보통 머리글 로고가 이미 있으므로 겹치지 않게 한다.
            header = doc.sections[0].header.paragraphs[0]
            header.add_run().add_picture(str(ctx.logo_path), height=Cm(1))
    _para(doc, ctx.title, "Title")
    _para(doc, f"벤더: {ctx.vendor}")
    _para(doc, f"언어: {language_native_name(ctx.target_language)}")
    _para(doc, f"문서 버전: {ctx.doc_version}")
    _para(doc, f"생성일: {ctx.created_at:%Y-%m-%d}")

    if ctx.revisions:
        _heading(doc, "문서 이력", 1)
        table = doc.add_table(rows=1, cols=len(REVISION_COLUMNS))
        _apply_style(table, "Table Grid")
        for cell, (_, label) in zip(table.rows[0].cells, REVISION_COLUMNS):
            cell.text = label
        for rev in ctx.revisions:
            row = table.add_row().cells
            for cell, (key, _) in zip(row, REVISION_COLUMNS):
                cell.text = str(rev.get(key) or "")

    doc.add_page_break()
    _heading(doc, "목차", 1)
    _add_toc_field(doc)
    doc.add_page_break()

    if ctx.overview:
        _heading(doc, "개요", 1)
        _add_markdown(doc, ctx.overview)

    for section in ctx.sections:
        _heading(doc, section.title, min(section.level, 4))
        if section.content:
            _add_markdown(doc, section.content)
        # 요약 본문에 이미 있는 주의사항은 다시 쓰지 않는다 (직접 수정 등으로 본문에서 빠진 것만)
        extra = [w for w in section.warnings if w not in (section.content or "")]
        if extra:
            p = doc.add_paragraph()
            p.add_run("주의 사항").bold = True
            for w in extra:
                _add_inline(_para(doc, "", "List Bullet"), w)
        label = page_label(section)
        if label:
            run = doc.add_paragraph().add_run(label)
            run.italic = True
            run.font.size = Pt(8)

    doc.save(str(out_path))


def _para(doc: DocxDocument, text: str, style: str | None = None) -> Paragraph:
    # 파트너 템플릿에 해당 스타일이 없을 수 있으므로 없으면 기본 스타일로 쓴다.
    try:
        return doc.add_paragraph(text, style=style)
    except KeyError:
        return doc.add_paragraph(text)


def _heading(doc: DocxDocument, text: str, level: int) -> None:
    try:
        doc.add_heading(text, level=level)
    except KeyError:
        doc.add_paragraph().add_run(text).bold = True


def _apply_style(obj, style: str) -> None:
    try:
        obj.style = style
    except (KeyError, ValueError):
        pass


def _clear_body(doc: DocxDocument) -> None:
    """양식 본문의 문단·표·그림을 모두 지운다. 마지막 구역 설정(용지, 여백, 머리글·바닥글 연결)은 남긴다."""
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _set_default_font(doc: DocxDocument) -> None:
    style = doc.styles["Normal"]
    style.font.size = Pt(10.5)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), KOREAN_FONT)


def _add_toc_field(doc: DocxDocument) -> None:
    """Word에서 열 때 갱신되는 목차 필드를 넣는다."""
    p = doc.add_paragraph()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), r'TOC \o "1-3" \h \z \u')
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "목차를 표시하려면 F9 키를 눌러 필드를 업데이트하세요."
    run.append(text)
    field.append(run)
    p._p.append(field)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)


def _add_markdown(doc: DocxDocument, markdown: str) -> None:
    in_code = False
    for line in markdown.splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            run = doc.add_paragraph().add_run(line)
            run.font.name = CODE_FONT
            run.font.size = Pt(9)
            continue
        if not line.strip():
            continue
        if m := _SUBHEADING.match(line):
            _add_inline(doc.add_paragraph(), f"**{m.group(1)}**")
        elif m := _BULLET.match(line):
            _add_inline(_para(doc, "", "List Bullet"), m.group(1))
        elif m := _NUMBERED.match(line):
            _add_inline(_para(doc, "", "List Number"), m.group(1))
        else:
            _add_inline(doc.add_paragraph(), line.lstrip("> ").strip())


def _add_inline(paragraph: Paragraph, text: str) -> None:
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            paragraph.add_run(part[1:-1]).font.name = CODE_FONT
        else:
            paragraph.add_run(part)
