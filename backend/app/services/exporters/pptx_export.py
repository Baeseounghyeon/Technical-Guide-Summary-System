import re
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.presentation import Presentation as PptxPresentation
from pptx.slide import SlideLayout
from pptx.util import Inches, Pt

from ..llm.base import language_native_name
from .base import REVISION_COLUMNS, ExportContext, open_template, page_label, strip_inline_markup

BULLETS_PER_SLIDE = 7
_LIST_PREFIX = re.compile(r"^\s*(?:[-*+]|#{1,6}|>)\s*")

_TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
_BODY_TYPES = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT}


def build_pptx(ctx: ExportContext, out_path: Path) -> None:
    if ctx.template_path:
        # 파트너 양식: 기존 슬라이드는 지우고 슬라이드 마스터·레이아웃(로고, 배경, 글꼴, 색, 크기)만 쓴다.
        prs = Presentation(open_template(ctx.template_path))
        _remove_all_slides(prs)
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    styled = ctx.template_path is not None  # 양식이면 글자 크기는 레이아웃에 맡긴다

    # 표지
    slide = prs.slides.add_slide(_title_layout(prs))
    _set_title(slide, ctx.title)
    subtitle = _placeholder(slide, {PP_PLACEHOLDER.SUBTITLE}) or _placeholder(slide, _BODY_TYPES)
    if subtitle is not None:
        subtitle.text = (
            f"{ctx.vendor} · {language_native_name(ctx.target_language)} · v{ctx.doc_version} · {ctx.created_at:%Y-%m-%d}"
        )
    if ctx.logo_path:
        slide.shapes.add_picture(str(ctx.logo_path), prs.slide_width - Inches(2.3), Inches(0.3), height=Inches(0.8))

    if ctx.revisions:
        _add_revision_slide(prs, ctx.revisions)

    if ctx.overview:
        _add_bullet_slides(prs, "개요", _markdown_lines(ctx.overview), styled)

    for section in ctx.sections:
        if not section.content:
            continue
        bullets = section.key_steps or _markdown_lines(section.content)
        bullets = bullets + [f"⚠ {w}" for w in section.warnings]
        label = page_label(section)
        if label:
            bullets.append(f"({label})")
        _add_bullet_slides(prs, section.title, bullets, styled)

    prs.save(str(out_path))


def _remove_all_slides(prs: PptxPresentation) -> None:
    slide_ids = prs.slides._sldIdLst
    for sld_id in list(slide_ids):
        prs.part.drop_rel(sld_id.rId)  # 관계가 끊긴 슬라이드는 저장할 때 빠진다
        slide_ids.remove(sld_id)


def _layout_types(layout: SlideLayout) -> set:
    return {ph.placeholder_format.type for ph in layout.placeholders}


def _title_layout(prs: PptxPresentation) -> SlideLayout:
    """표지 레이아웃: 가운데 제목 또는 제목+부제목이 있는 레이아웃. 없으면 첫 레이아웃."""
    for layout in prs.slide_layouts:
        types = _layout_types(layout)
        if PP_PLACEHOLDER.CENTER_TITLE in types or {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.SUBTITLE} <= types:
            return layout
    return prs.slide_layouts[0]


def _content_layout(prs: PptxPresentation) -> SlideLayout:
    """본문 레이아웃: 제목과 본문 영역이 있는 레이아웃 (양식마다 순서·이름이 달라 자리표시자로 찾는다)."""
    for layout in prs.slide_layouts:
        types = _layout_types(layout)
        if types & _TITLE_TYPES and types & _BODY_TYPES:
            return layout
    for layout in prs.slide_layouts:
        if _layout_types(layout) & _TITLE_TYPES:
            return layout
    return prs.slide_layouts[0]


def _placeholder(slide, types: set):
    return next((ph for ph in slide.placeholders if ph.placeholder_format.type in types), None)


def _set_title(slide, text: str) -> None:
    title = slide.shapes.title or _placeholder(slide, _TITLE_TYPES)
    if title is not None:
        title.text = text


def _markdown_lines(markdown: str) -> list[str]:
    lines = []
    for line in markdown.splitlines():
        if not line.strip() or line.strip().startswith("```"):
            continue
        lines.append(strip_inline_markup(_LIST_PREFIX.sub("", line).strip()))
    return lines


def _add_bullet_slides(prs: PptxPresentation, title: str, bullets: list[str], styled: bool) -> None:
    chunks = [bullets[i : i + BULLETS_PER_SLIDE] for i in range(0, len(bullets), BULLETS_PER_SLIDE)] or [[]]
    for index, chunk in enumerate(chunks):
        slide = prs.slides.add_slide(_content_layout(prs))
        _set_title(slide, title if index == 0 else f"{title} (계속)")
        body = _placeholder(slide, _BODY_TYPES)
        if body is None:
            body = slide.shapes.add_textbox(Inches(0.7), Inches(1.6), prs.slide_width - Inches(1.4), prs.slide_height - Inches(2.2))
        frame = body.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE  # 넘치면 PowerPoint 가 글자를 줄인다
        for i, text in enumerate(chunk):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.text = text
            if not styled:
                paragraph.font.size = Pt(18)


def _add_revision_slide(prs: PptxPresentation, revisions: list[dict]) -> None:
    slide = prs.slides.add_slide(_content_layout(prs))
    _set_title(slide, "문서 이력")
    body = _placeholder(slide, _BODY_TYPES)
    if body is not None:
        body._element.getparent().remove(body._element)
    rows, cols = len(revisions) + 1, len(REVISION_COLUMNS)
    table = slide.shapes.add_table(
        rows, cols, Inches(0.5), Inches(1.6), prs.slide_width - Inches(1), Inches(0.4) * rows
    ).table
    for c, (_, label) in enumerate(REVISION_COLUMNS):
        table.cell(0, c).text = label
    for r, rev in enumerate(revisions, start=1):
        for c, (key, _) in enumerate(REVISION_COLUMNS):
            table.cell(r, c).text = str(rev.get(key) or "")
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(11)
