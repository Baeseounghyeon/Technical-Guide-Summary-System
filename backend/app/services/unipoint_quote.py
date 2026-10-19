"""유니포인트 견적서(엑셀) 만들기 — 견적서_매핑정리.md 를 그대로 따른다.

양식: 받은 견적서 엑셀에서 건별 값과 계산 메모만 지운 것 (로고·직인·서식·인쇄 영역 A1:J 유지).
      직인 그림이 있어 공개 저장소(git)에 넣지 않는다 → 앱 저장소(Supabase, 비공개)의 assets/quote_unipoint.xlsx.
      개발 PC 에서는 app/assets/quote_unipoint.xlsx (git 제외) 가 있으면 그것을 쓴다.

견적서 한 행 = 매입 견적 PDF 한 건 (매핑 1·2절)
    B  Product Number      = 대표 라인의 Maintenance Part Number
    C  Product Name        = 대표 라인 Covered Product 의 제품군 (예: 'Luna Network HSM')
    D  Product Description = 대상 제품(대문자, GRK-16→GRK16) / Serial Num. ○ / ○ / (시작 ~ 종료) 1year MA
    E  수량                = 대표 라인 수량 (정수)
    F  단위                = 식
    J  비고                = 'Enhanced MA' + (계약 번호)
    N  매입단가(USD)       = 포함한 라인의 Net Line Total 합 (여러 라인이면 Quote Total)
    O  환율 · P 수량(1) · M =N*P · Q 매입 =N*O*P
    G  금액 =N*O*2.3 · H 공급가 =G*(1-I) · I 할인율(값) · R GP =H-Q · S GP율 =R/H
    합계 H =SUM, 부가세 포함 =H*1.1, B16 =NUMBERSTRING 제안금액, P9 목표 공급가 · P10 =1-P9/SUM(G) (할인율 역산)
"""

import html
import io
import re
from contextlib import contextmanager
from copy import copy
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import pymupdf
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_boundaries

from .thales_pdf import PURCHASE_COLUMNS

TEMPLATE = Path(__file__).resolve().parents[1] / "assets" / "quote_unipoint.xlsx"
TEMPLATE_KEY = "assets/quote_unipoint.xlsx"


@contextmanager
def template_path():
    """견적서 양식 파일 경로: 개발 PC 의 파일, 없으면 앱 저장소에서 받는다."""
    if TEMPLATE.exists():
        yield TEMPLATE
        return
    from .storage import get_storage

    storage = get_storage()
    if not storage.exists(TEMPLATE_KEY):
        raise RuntimeError("견적서 기본 양식(assets/quote_unipoint.xlsx)이 저장소에 없습니다. 관리자에게 문의하세요.")
    with storage.local_path(TEMPLATE_KEY) as path:
        yield path
MEDIA_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass
class QuoteSettings:
    exchange_rate: float = 1360.0  # O열 환율
    multiplier: float = 2.3  # G = N × O × 배수
    discount_pct: float = 0.0  # 할인율(%) — 목표 금액이 없을 때
    target_total: float | None = None  # P9 목표 공급가 합계. 있으면 I = 1 - P9 / SUM(G)


@dataclass
class Purchase:
    """매입 견적 PDF 한 건 (화면에서 고친 값)."""

    rows: list[list[str]]  # PURCHASE_COLUMNS 순서
    included: list[bool]  # 금액(N)에 넣을 라인
    main: int  # 대표 라인 (품명·설명·시리얼·수량·계약 번호)
    meta: dict = field(default_factory=dict)  # quote_number, end_user, partner, quote_total


def _v(row: list[str], name: str) -> str:
    i = PURCHASE_COLUMNS.index(name)
    return (row[i] if i < len(row) else "").strip()


def _num(text) -> float:
    try:
        return float(str(text).replace(",", "").strip() or 0)
    except ValueError:
        return 0.0


def default_main(rows: list[list[str]], common_parts: dict[str, int] | None = None) -> int:
    """대표 라인 기본값: 여러 매입 견적에 공통으로 나오는 대상 제품(예: S750) 중 수량이 가장 큰 라인."""
    if not rows:
        return 0
    common_parts = common_parts or {}
    return max(range(len(rows)), key=lambda i: (common_parts.get(_v(rows[i], "대상 제품 번호"), 0), _num(_v(rows[i], "수량")), -i))


def product_family(desc: str) -> str:
    """'LUNA NETWORK HSM S750 (PED, …)' → 'Luna Network HSM' (모델 번호 앞까지)."""
    head = desc.split("(")[0].strip()
    words = head.split()
    if words and re.search(r"\d", words[-1]):
        words = words[:-1]
    small = {"HSM", "PED", "USB", "PCIE", "CTM", "CTE", "DDC"}
    return " ".join(w if w.upper() in small else w.capitalize() for w in words) or head


def _description(row: list[str]) -> str:
    desc = _v(row, "대상 제품").upper()
    head, _, rest = desc.partition("(")
    spec = ("(" + rest) if rest else ""
    spec = re.sub(r",\s+", ",", spec).replace("GRK-16", "GRK16")
    lines = [head.strip() + " "]  # 원래 견적서처럼 모델명 뒤에 한 칸
    if spec:
        lines.append(spec)
    serial = _v(row, "시리얼 번호")
    if serial:
        lines.append(f"Serial Num. {serial}")
    months = re.search(r"(\d+)", _v(row, "커버리지"))
    ma = f"{int(months.group(1)) // 12}year MA" if months and int(months.group(1)) % 12 == 0 else (f"{_v(row, '커버리지')} MA" if _v(row, "커버리지") else "")
    start, end = _v(row, "시작일"), _v(row, "종료일")
    if start or end:
        lines.append(f"({start} ~ {end}) {ma}".strip())
    elif ma:
        lines.append(ma)
    return "\n".join(lines)


def _service(text: str) -> str:
    low = text.lower()
    for key, label in (("premier", "Premier MA"), ("enhanced", "Enhanced MA"), ("standard", "Standard MA")):
        if key in low:
            return label
    return "MA" if "maintenance" in low else text


def _period_days(row: list[str]) -> float:
    """기간 길이(일). 시작·종료일이 없으면 커버리지 개월 수로 어림한다."""
    try:
        start = datetime.strptime(_v(row, "시작일"), "%Y.%m.%d")
        end = datetime.strptime(_v(row, "종료일"), "%Y.%m.%d")
        return (end - start).days
    except ValueError:
        months = re.search(r"(\d+)", _v(row, "커버리지"))
        return int(months.group(1)) * 30.44 if months else 0


def _line_order(line: dict) -> tuple:
    """품목 순서: 수량 오름차순 → 수량이 같으면 기간이 긴 것이 뒤 → 기간도 같으면 Enhanced MA(계약 번호)가 큰 것이 뒤."""
    digits = re.findall(r"\d+", line.get("contract") or "")
    return (line["quantity"], line["period_days"], int(digits[-1]) if digits else 0)


def build_lines(purchases: list[Purchase], s: QuoteSettings) -> tuple[list[dict], dict]:
    """매입 견적 PDF 마다 견적서 한 행. 금액·공급가·매입·GP 도 엑셀 수식과 같은 방식으로 계산한다."""
    lines = []
    for p in purchases:
        if not p.rows:
            continue
        main = p.rows[min(max(p.main, 0), len(p.rows) - 1)]
        cost = sum(_num(_v(r, "합계(USD)")) for r, inc in zip(p.rows, p.included) if inc)
        contract = _v(main, "계약 번호")
        lines.append({
            "product_number": _v(main, "유지보수 제품 번호") or _v(main, "대상 제품 번호"),
            "product_name": product_family(_v(main, "대상 제품")),
            "description": _description(main),
            "quantity": int(_num(_v(main, "수량"))),
            "unit": "식",
            "note": "\n".join(x for x in (_service(_v(main, "서비스")), f"({contract})" if contract else "") if x),
            "cost_usd": round(cost, 2),  # N
            "quote_number": p.meta.get("quote_number", ""),
            "period_days": _period_days(main),
            "contract": contract,
        })
    lines.sort(key=_line_order)
    rate, mult = s.exchange_rate, s.multiplier
    amount_total = sum(line["cost_usd"] * rate * mult for line in lines)
    if s.target_total and amount_total:
        discount = 1 - s.target_total / amount_total
    else:
        discount = min(max(s.discount_pct, 0.0), 100.0) / 100
    for i, line in enumerate(lines, start=1):
        line["no"] = i
        line["amount"] = line["cost_usd"] * rate * mult  # G
        line["discount"] = discount  # I
        line["supply"] = line["amount"] * (1 - discount)  # H
        line["cost_krw"] = line["cost_usd"] * rate  # Q
        line["gp"] = line["supply"] - line["cost_krw"]  # R
        line["gp_rate"] = line["gp"] / line["supply"] if line["supply"] else 0  # S
    supply = sum(line["supply"] for line in lines)
    cost_krw = sum(line["cost_krw"] for line in lines)
    totals = {
        "amount": amount_total,
        "discount": discount,
        "supply": supply,
        "vat_included": supply * 1.1,
        "cost_krw": cost_krw,
        "gp": supply - cost_krw,
        "gp_rate": (supply - cost_krw) / supply if supply else 0,
    }
    return lines, totals


_DIGITS = "영일이삼사오육칠팔구"


def korean_amount(n: float) -> str:
    """엑셀 NUMBERSTRING(n,1) 과 같은 읽기: 180180000 → '일억팔천일십팔만'."""
    n = int(round(n))
    if n <= 0:
        return "영"
    out = []
    for gi, unit_name in enumerate(["", "만", "억", "조"]):
        part = (n // 10000**gi) % 10000
        if part:
            text = "".join(_DIGITS[part // v % 10] + u for u, v in (("천", 1000), ("백", 100), ("십", 10), ("", 1)) if part // v % 10)
            out.append(text + unit_name)
    return "".join(reversed(out))


@dataclass
class QuoteDoc:
    project: str  # B8 프로젝트명
    recipient: str  # B7 ○○ 귀중
    date: str  # H8 견적일자 (YYYY-MM-DD)
    author: dict[str, str]  # B2 부서 / 이름 / 전화 / 메일
    lines: list[dict]
    totals: dict
    settings: QuoteSettings


# ---------- 엑셀 ----------


def _shift_rows(ws, at: int, count: int) -> None:
    """at 행 앞에 count 개 행을 넣는다. openpyxl 이 옮기지 않는 병합 칸·행 높이·그림·인쇄 영역도 함께 옮긴다."""
    if count <= 0:
        return
    heights = {r: d.height for r, d in ws.row_dimensions.items() if d.height is not None and r >= at}
    merged = [m for m in ws.merged_cells.ranges if m.min_row >= at]
    for m in merged:
        ws.merged_cells.remove(m)
    ws.insert_rows(at, count)
    for m in merged:
        m.shift(0, count)
        ws.merged_cells.add(m)
    for r in heights:
        ws.row_dimensions[r].height = None
    for r, h in heights.items():
        ws.row_dimensions[r + count].height = h
    for image in ws._images:
        anchor = image.anchor
        if hasattr(anchor, "_from") and anchor._from.row >= at - 1:
            anchor._from.row += count
            if getattr(anchor, "to", None) is not None:
                anchor.to.row += count
    if ws.print_area:
        ref = (ws.print_area if isinstance(ws.print_area, str) else ws.print_area[0]).split("!")[-1].replace("$", "")
        min_col, min_row, max_col, max_row = range_boundaries(ref)
        ws.print_area = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row + count}"


def author_line(a: dict) -> str:
    """B2 담당자 줄: '부서/ 이름 / 전화 / 메일' (원래 견적서 표기)."""
    head = f"{a.get('department', '')}/ {a.get('name', '')}" if a.get("department") else a.get("name", "")
    return " / ".join(x for x in (head, a.get("phone", ""), a.get("email", "")) if x)


def _find(ws, test, rows=None):
    """조건에 맞는 첫 칸 (행 순서)."""
    for row in ws.iter_rows(min_row=rows[0] if rows else 1, max_row=rows[1] if rows else ws.max_row):
        for c in row:
            if isinstance(c.value, str) and test(c.value):
                return c
    return None


@dataclass
class FormLayout:
    """유니포인트 견적서 양식의 위치 (머리글 글자로 찾는다)."""

    header: int  # 'Product Number' 머리글 행
    first: int  # 첫 품목 행
    blank: int  # '- 이하여백 -' 행 (품목 행은 first ~ blank-1)
    total: int  # 제안가 합계 (부가세별도)
    vat: int  # 제안가 합계 (부가세포함)
    cols: dict  # 항목 → 열 문자
    proposal: str  # 제안금액 칸
    recipient: str
    project: str
    date: str
    target: str  # 목표 공급가 (P9)
    discount: str  # 할인율 역산 (P10)


def find_layout(ws) -> FormLayout | None:
    """유니포인트 견적서 모양(머리글 'Product Number' · '- 이하여백 -' · 제안가 합계)이 아니면 None."""
    head = _find(ws, lambda v: v.strip().lower() == "product number")
    if head is None:
        return None
    h = head.row
    blank = _find(ws, lambda v: "이하여백" in v, (h + 1, ws.max_row))
    if blank is None:
        return None
    total = _find(ws, lambda v: "부가세별도" in v.replace(" ", ""), (blank.row, ws.max_row))
    vat = _find(ws, lambda v: "부가세포함" in v.replace(" ", ""), (blank.row, ws.max_row))
    if total is None or vat is None:
        return None
    names: dict[str, list[str]] = {}
    for c in ws[h]:
        if isinstance(c.value, str) and c.value.strip():
            names.setdefault(c.value.strip(), []).append(c.column_letter)
    pick = lambda name, default, nth=0: (names.get(name) or [])[nth] if len(names.get(name) or []) > nth else default  # noqa: E731
    pn_col = head.column
    cols = {
        "no": get_column_letter(max(pn_col - 1, 1)),
        "product_number": head.column_letter,
        "product_name": pick("Product Name", "C"),
        "description": pick("Product Description", "D"),
        "quantity": pick("수량", "E"),
        "unit": pick("단위", "F"),
        "amount": pick("금액", "G"),
        "supply": pick("공급가", "H"),
        "discount": pick("할인율", "I"),
        "note": pick("비고", "J"),
        "cost": pick("매입단가", "N"),
        "rate": pick("환율", "O"),
        "cost_qty": pick("수량", "P", 1),
        "cost_krw": pick("매입", "Q"),
        "gp": pick("GP", "R"),
    }
    cols["cost_m"] = get_column_letter(max(column_index_from_string(cols["cost"]) - 1, 1))  # M = N*P
    cols["gp_rate"] = pick("GP율", get_column_letter(column_index_from_string(cols["gp"]) + 1))
    proposal = _find(ws, lambda v: "제안금액" in v or "NUMBERSTRING" in v.upper(), (1, h))
    recipient = _find(ws, lambda v: "귀중" in v, (1, h))
    project = _find(ws, lambda v: "프로젝트명" in v, (1, h))
    date_label = _find(ws, lambda v: v.replace(" ", "").startswith("견적일자"), (1, h))
    target_label = _find(ws, lambda v: "목표금액" in v.replace(" ", ""), (1, h))
    return FormLayout(
        header=h,
        first=h + 1,
        blank=blank.row,
        total=total.row,
        vat=vat.row,
        cols=cols,
        proposal=proposal.coordinate if proposal else f"B{h - 1}",
        recipient=recipient.coordinate if recipient else "B7",
        project=project.coordinate if project else "B8",
        date=f"{get_column_letter(date_label.column + 1)}{date_label.row}" if date_label else "H8",
        target=f"{target_label.column_letter}{target_label.row + 1}" if target_label else "P9",
        discount=f"{target_label.column_letter}{target_label.row + 2}" if target_label else "P10",
    )


def is_quote_form(path: Path) -> bool:
    """올린 엑셀 양식이 유니포인트 견적서 모양이면 표 안에 바로 채운다."""
    try:
        wb = load_workbook(str(path))
    except Exception:  # noqa: BLE001
        return False
    return find_layout(wb.worksheets[0]) is not None


def build_xlsx(q: QuoteDoc, template: Path | None = None) -> tuple[bytes, dict[str, object]]:
    """양식의 품목 표(머리글 아래 ~ '- 이하여백 -' 위) 안에 값을 채운다. 품목이 표보다 많을 때만 행을 늘린다.
    (견적서 엑셀, 미리보기용 수식 칸 계산값 {칸: 값})"""
    if template is None:
        with template_path() as path:
            wb = load_workbook(str(path))
    else:
        wb = load_workbook(str(template))
    ws = wb.worksheets[0]
    L = find_layout(ws)
    if L is None:
        raise ValueError("견적서 양식에서 품목 표('Product Number' 머리글, '- 이하여백 -', 제안가 합계)를 찾지 못했습니다.")
    capacity = L.blank - L.first
    extra = len(q.lines) - capacity
    if extra > 0:  # 표보다 품목이 많으면 '- 이하여백 -' 위에 행을 늘린다 (마지막 품목 행 서식으로)
        _shift_rows(ws, L.blank, extra)
        src = L.blank - 1
        for r in range(L.blank, L.blank + extra):
            for col in range(1, ws.max_column + 1):
                a = ws.cell(row=src, column=col)
                if a.has_style:
                    ws.cell(row=r, column=col)._style = copy(a._style)
            ws.row_dimensions[r].height = ws.row_dimensions[src].height
        L = find_layout(ws)
    c = L.cols
    s = q.settings
    calc: dict[str, object] = {}
    a = q.author
    # 회사명·주소·연락처·담당자 줄은 양식에서 비워 두고 웹의 작성자 정보로 채운다
    for ref, key in (("B1", "company"), ("J1", "address"), ("J2", "contact")):
        if (a.get(key) or "").strip():
            ws[ref] = a[key].strip()
    ws["B2"] = author_line(a)  # '주식회사 ○○' 아래 줄: 작성자 정보
    ws[L.recipient] = f"{q.recipient} 귀중"
    ws[L.project] = f"프로젝트명 : {q.project}"
    ws[L.date] = datetime.strptime(q.date, "%Y-%m-%d")
    last = L.first + max(len(q.lines), 1) - 1
    if s.target_total:  # 할인율 역산 (원래 견적서 O8~P10)
        ws[L.target] = s.target_total
        ws[L.discount] = f"=1-{L.target}/SUM({c['amount']}{L.first}:{c['amount']}{last})"
        calc[L.discount] = q.totals["discount"]
    else:
        ws[L.target] = None
        ws[L.discount] = None
    for r in range(L.first, L.blank):  # 표 안의 이전 값을 비운 뒤 채운다
        for key in c:
            ws[f"{c[key]}{r}"] = None
    for i, line in enumerate(q.lines):
        r = L.first + i
        cell = lambda key: f"{c[key]}{r}"  # noqa: E731
        ws[cell("no")] = line["no"]
        ws[cell("product_number")] = line["product_number"]
        ws[cell("product_name")] = line["product_name"]
        ws[cell("description")] = line["description"]
        ws[cell("quantity")] = line["quantity"]
        ws[cell("unit")] = line["unit"]
        ws[cell("amount")] = f"={cell('cost')}*{cell('rate')}*{s.multiplier:g}"
        ws[cell("supply")] = f"={cell('amount')}*(1-{cell('discount')})"
        ws[cell("discount")] = line["discount"]  # 원래 견적서처럼 역산한 값을 숫자로
        ws[cell("note")] = line["note"]
        ws[cell("cost_m")] = f"={cell('cost')}*{cell('cost_qty')}"
        ws[cell("cost")] = line["cost_usd"]
        ws[cell("rate")] = s.exchange_rate
        ws[cell("cost_qty")] = 1
        ws[cell("cost_krw")] = f"={cell('cost')}*{cell('rate')}*{cell('cost_qty')}"
        ws[cell("gp")] = f"={cell('supply')}-{cell('cost_krw')}"
        ws[cell("gp_rate")] = f"={cell('gp')}/{cell('supply')}"
        calc.update({cell("amount"): line["amount"], cell("supply"): line["supply"], cell("cost_m"): line["cost_usd"],
                     cell("cost_krw"): line["cost_krw"], cell("gp"): line["gp"], cell("gp_rate"): line["gp_rate"]})
    h, q_, r_, s_ = c["supply"], c["cost_krw"], c["gp"], c["gp_rate"]
    ws[f"{h}{L.total}"] = f"=SUM({h}{L.first}:{h}{L.blank})"
    ws[f"{q_}{L.total}"] = f"=SUM({q_}{L.first}:{q_}{L.blank})"
    ws[f"{r_}{L.total}"] = f"={h}{L.total}-{q_}{L.total}"
    ws[f"{s_}{L.total}"] = f"={r_}{L.total}/{h}{L.total}"
    ws[f"{h}{L.vat}"] = f"={h}{L.total}*1.1"
    ws[L.proposal] = f'="제안금액 : 일금"&NUMBERSTRING({h}{L.vat},1)&"원정(부가세 포함)"'
    t = q.totals
    calc.update({f"{h}{L.total}": t["supply"], f"{q_}{L.total}": t["cost_krw"], f"{r_}{L.total}": t["gp"], f"{s_}{L.total}": t["gp_rate"],
                 f"{h}{L.vat}": t["vat_included"], L.proposal: f"제안금액 : 일금{korean_amount(t['vat_included'])}원정(부가세 포함)"})
    wb.calculation.fullCalcOnLoad = True  # 엑셀로 열면 모든 수식을 다시 계산한다
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), calc


# ---------- 미리보기 (인쇄 영역 A1:J) ----------


def _fmt(value, number_format: str) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return f"{value.year}년 {value.month}월 {value.day}일"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if "%" in number_format:
            return f"{value * 100:.1f}%"
        if "#,##0" in number_format:
            return f"{value:,.0f}"
        return f"{value:g}"
    return str(value)


def preview_html(data: bytes, calc: dict) -> str:
    """고객에게 나가는 인쇄 영역만 그린다 (병합·굵기·채움·테두리·정렬). 수식 칸은 계산값."""
    ws = load_workbook(io.BytesIO(data)).worksheets[0]
    ref = (ws.print_area if isinstance(ws.print_area, str) else ws.print_area[0]).split("!")[-1].replace("$", "")
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    spans, covered = {}, set()
    for m in ws.merged_cells.ranges:
        spans[(m.min_row, m.min_col)] = (m.max_row - m.min_row + 1, m.max_col - m.min_col + 1)
        covered |= {(r, c) for r in range(m.min_row, m.max_row + 1) for c in range(m.min_col, m.max_col + 1)} - {(m.min_row, m.min_col)}
    cols = "".join(f'<col style="width:{round((ws.column_dimensions[get_column_letter(c)].width or 8.43) * 7)}px">' for c in range(min_col, max_col + 1))
    def empty(r: int, c: int) -> bool:
        v = calc.get(f"{get_column_letter(c)}{r}", ws.cell(row=r, column=c).value)
        return v in (None, "") and (r, c) not in covered and (r, c) not in spans

    rows = []
    for r in range(min_row, max_row + 1):
        cells = []
        spill_until = 0
        for c in range(min_col, max_col + 1):
            if (r, c) in covered or c <= spill_until:
                continue
            cell = ws.cell(row=r, column=c)
            value = calc.get(cell.coordinate, cell.value)
            # 엑셀처럼 줄바꿈 없는 글은 오른쪽 빈 칸으로 넘쳐 보이게 한다
            if (r, c) not in spans and isinstance(value, str) and value and not (cell.alignment and cell.alignment.wrap_text):
                end = c
                while end + 1 <= max_col and empty(r, end + 1):
                    end += 1
                if end > c:
                    spans[(r, c)] = (1, end - c + 1)
                    spill_until = end
            if isinstance(value, str) and value.startswith("="):
                value = ""
            style = []
            if cell.font and cell.font.b:
                style.append("font-weight:bold")
            if cell.font and cell.font.sz:
                style.append(f"font-size:{round(cell.font.sz * 1.2)}px")
            fg = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type == "solid" else None
            if isinstance(fg, str) and fg not in ("00000000", "FFFFFFFF"):
                style.append(f"background:#{fg[-6:]}")
            for side in ("left", "right", "top", "bottom"):
                b = getattr(cell.border, side)
                if b is not None and b.style:
                    style.append(f"border-{side}:1px solid #444")
            h = cell.alignment.horizontal if cell.alignment else None
            if h:
                style.append(f"text-align:{'center' if 'center' in h else h}")
            rs, cs = spans.get((r, c), (1, 1))
            attrs = (f' rowspan="{rs}"' if rs > 1 else "") + (f' colspan="{cs}"' if cs > 1 else "")
            cells.append(f'<td{attrs} style="{";".join(style)}">{html.escape(_fmt(value, cell.number_format or ""))}</td>')
        height = ws.row_dimensions[r].height
        rows.append(f'<tr style="height:{round((height or 15) * 1.33)}px">{"".join(cells)}</tr>')
    return (
        "<style>body{font-family:'Malgun Gothic',sans-serif;color:#111;margin:12px;font-size:12px}"
        "table{border-collapse:collapse;table-layout:fixed}td{padding:2px 4px;white-space:pre-wrap;vertical-align:middle}</style>"
        f"<table><colgroup>{cols}</colgroup>{''.join(rows)}</table>"
        "<p style='color:#777;font-size:11px'>※ 로고·직인 그림은 미리보기에 나오지 않습니다(엑셀 파일에는 들어 있습니다). 매입·GP 칸(인쇄 영역 밖)도 엑셀 파일에만 있습니다.</p>"
    )


# ---------- PDF (고객용 A~J 열 내용, 엑셀 견적서와 같은 순서) ----------

def build_pdf(q: QuoteDoc, template: Path | None = None) -> bytes:
    """고객에게 보내는 내용(인쇄 영역 A~J)만 담는다. 매입·GP 는 넣지 않는다."""
    if template is None:
        with template_path() as path:
            wb = load_workbook(str(path))
    else:
        wb = load_workbook(str(template))
    ws = wb.worksheets[0]
    esc = html.escape
    fixed = lambda ref: esc(str(ws[ref].value or ""))  # noqa: E731 — 양식의 고정 문구
    # 양식에 남아 있는 Remark 문구 (합계 행 아래 B열)
    layout = find_layout(ws)
    start = (layout.vat + 1) if layout else ws.max_row + 1
    remarks = [ws.cell(row=r, column=2).value for r in range(start, ws.max_row + 1) if ws.cell(row=r, column=2).value]
    a = q.author
    th = "border:1px solid #666;background-color:#DCE3EE;padding:3px;font-size:7.5pt;text-align:center;font-weight:bold"
    td = "border:1px solid #666;padding:3px;font-size:7.5pt;vertical-align:middle"
    cols = [("no", "", "center"), ("product_number", "Product Number", "center"), ("product_name", "Product Name", "center"),
            ("description", "Product Description", "left"), ("quantity", "수량", "center"), ("unit", "단위", "center"),
            ("amount", "금액", "right"), ("supply", "공급가", "right"), ("discount", "할인율", "center"), ("note", "비고", "center")]

    def text(line, key):
        v = line[key]
        if key in ("amount", "supply"):
            return f"{v:,.0f}"
        if key == "discount":
            return f"{v * 100:.1f}%"
        return esc(str(v)).replace(chr(10), "<br/>")

    head = "".join(f'<th style="{th}">{esc(label)}</th>' for _, label, _ in cols)
    body = "".join(
        "<tr>" + "".join(f'<td style="{td};text-align:{align}">{text(line, key)}</td>' for key, _, align in cols) + "</tr>" for line in q.lines
    )
    body += f'<tr><td colspan="10" style="{td};text-align:center;color:#555">- 이하여백 -</td></tr>'
    for label, value in (("제안가 합계 (부가세별도)", q.totals["supply"]), ("제안가 합계 (부가세포함)", q.totals["vat_included"])):
        body += (f'<tr><td colspan="7" style="{td};text-align:center;font-weight:bold">{label}</td>'
                 f'<td style="{td};text-align:right;font-weight:bold">{value:,.0f}</td><td colspan="2" style="{td}"></td></tr>')
    date = datetime.strptime(q.date, "%Y-%m-%d")
    terms = "".join(f'<p style="margin:1px 0;font-size:9pt">{fixed(ref)}</p>' for ref in ("B10", "B11", "B12", "B13", "B14") if ws[ref].value)
    # 회사 정보: 웹에서 입력한 값, 없으면 양식에 적힌 값
    company = (a.get("company") or "").strip() or str(ws["B1"].value or "")
    address = (a.get("address") or "").strip() or str(ws["J1"].value or "")
    contact = (a.get("contact") or "").strip() or str(ws["J2"].value or "")
    doc_html = (
        f'<p style="font-size:10pt;font-weight:bold;margin:0">{esc(company)}</p>'
        f'<p style="font-size:8.5pt;margin:0">{esc(author_line(a))}</p>'
        f'<p style="font-size:7.5pt;color:#444;margin:0">{esc(" · ".join(x for x in (address, contact) if x))}</p>'
        '<p style="text-align:center;font-size:22pt;font-weight:bold;margin:8px 0 10px 0">견 적 서</p>'
        f'<p style="font-size:13pt;font-weight:bold;margin:0 0 4px 0">{esc(q.recipient)} 귀중</p>'
        f'<p style="font-size:10pt;margin:0">프로젝트명 : {esc(q.project)}</p>'
        f'<p style="font-size:9pt;margin:2px 0 6px 0">견 적 일 자 : {date.year}년 {date.month}월 {date.day}일</p>'
        f"{terms}"
        f'<p style="font-size:11pt;font-weight:bold;margin:8px 0 4px 0">제안금액 : 일금{korean_amount(q.totals["vat_included"])}원정(부가세 포함)'
        '<span style="font-size:8pt;font-weight:normal">  ( 단위 : 원 )</span></p>'
        f'<table style="width:100%;border-collapse:collapse"><tr>{head}</tr>{body}</table>'
        + "".join(f'<p style="font-size:8pt;margin:1px 0">{esc(str(t))}</p>' for t in remarks)
    )
    story = pymupdf.Story(html=doc_html, user_css="body{font-family:sans-serif}")
    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    page = pymupdf.paper_rect("a4-l")
    more = True
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(page + (30, 30, -30, -30))
        story.draw(device)
        writer.end_page()
    writer.close()
    with pymupdf.open("pdf", buffer.getvalue()) as pdf:
        pdf.subset_fonts()
        return pdf.tobytes(garbage=4, deflate=True)
