import re
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pptx import Presentation

from .base import FRONT_MATTER_TITLE, ParsedDocument, ParsedSection

_HEADING_STYLE = re.compile(r"^(?:Heading|제목)\s*(\d)$", re.IGNORECASE)


def parse_docx(path: Path) -> ParsedDocument:
    doc = Document(str(path))
    sections: list[ParsedSection] = []
    current: ParsedSection | None = None

    def append(text: str) -> None:
        nonlocal current
        if current is None:
            current = ParsedSection(FRONT_MATTER_TITLE, 1)
            sections.append(current)
        current.text += text + "\n"

    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            level = _heading_level(block)
            if level:
                current = ParsedSection(text, level)
                sections.append(current)
            else:
                append(text)
        elif isinstance(block, Table):
            for row in block.rows:
                append(" | ".join(cell.text.strip() for cell in row.cells))

    for s in sections:
        s.text = s.text.strip()
    return ParsedDocument(title=doc.core_properties.title or None, page_count=None, sections=sections)


def _heading_level(paragraph: Paragraph) -> int | None:
    name = paragraph.style.name if paragraph.style is not None else ""
    m = _HEADING_STYLE.match(name)
    return int(m.group(1)) if m else None


def parse_pptx(path: Path) -> ParsedDocument:
    prs = Presentation(str(path))
    sections: list[ParsedSection] = []
    for index, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None else ""
        texts: list[str] = []
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(cell.text.strip() for cell in row.cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append(f"[발표자 노트] {notes}")
        body = "\n".join(texts)

        # 제목이 같은 연속 슬라이드는 한 섹션으로 묶는다.
        if sections and title and sections[-1].title == title:
            sections[-1].text += "\n" + body
            sections[-1].page_end = index
        else:
            sections.append(ParsedSection(title or f"슬라이드 {index}", 1, body, index, index))

    return ParsedDocument(title=prs.core_properties.title or None, page_count=len(prs.slides), sections=sections)


def sheet_rows(path: Path) -> tuple[str | None, list[tuple[str, list[list]]]]:
    """엑셀 시트별 행 (제목, [(시트 이름, 행들)]). .xlsx 는 openpyxl, 예전 형식 .xls 는 xlrd 로 읽는다."""
    if path.suffix.lower() == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path))
        sheets = []
        for sh in book.sheets():
            rows = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                    elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        row.append(None)
                    else:
                        row.append(cell.value)
                rows.append(row)
            sheets.append((sh.name, rows))
        return None, sheets
    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        title = wb.properties.title if wb.properties else None
        return title, [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    finally:
        wb.close()


def parse_xlsx(path: Path) -> ParsedDocument:
    """시트마다 한 섹션(행은 ' | ' 로 이은 줄). '항목 | 값' 처럼 두 칸만 찬 행은 그 아래 하위 섹션으로도 둔다
    (규격서 재사용 때 항목 이름으로 값을 찾기 위해). .xls 도 같은 방식."""
    title, sheets = sheet_rows(path)
    sections: list[ParsedSection] = []
    for index, (sheet_name, rows) in enumerate(sheets, start=1):
        lines: list[str] = []
        pairs: list[ParsedSection] = []
        for row in rows:
            cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if not cells:
                continue
            lines.append(" | ".join(cells))
            if len(cells) == 2:
                pairs.append(ParsedSection(cells[0][:500], 2, cells[1], index, index))
        sections.append(ParsedSection(sheet_name or f"시트 {index}", 1, "\n".join(lines), index, index))
        sections += pairs
    return ParsedDocument(title=title or None, page_count=len(sheets), sections=sections)
