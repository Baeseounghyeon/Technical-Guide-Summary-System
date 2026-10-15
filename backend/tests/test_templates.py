"""파트너 양식(Word/PPT 템플릿): 양식의 기존 내용은 지우고 디자인(머리글 로고, 스타일, 마스터, 크기)만 쓴다."""

import io
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Cm
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.worker import run_pending_jobs

from .conftest import ADMIN, login
from .test_flow import _make_vendor_docx, _partner_with_vendor, _signup_and_approve


def _docx_template(path: Path, logo: Path) -> None:
    d = Document()
    header = d.sections[0].header.paragraphs[0]
    header.text = "파트너 머리글 "
    header.add_run().add_picture(str(logo), height=Cm(1))
    d.add_paragraph("기존 본문 - 지워져야 함")
    d.add_table(rows=1, cols=1).cell(0, 0).text = "기존 표 - 지워져야 함"
    d.save(str(path))


def _to_dotx(src: Path, dst: Path) -> None:
    """docx 를 Word 서식 파일(.dotx) 형식으로 바꾼다."""
    buf = io.BytesIO()
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"wordprocessingml.document.main+xml", b"wordprocessingml.template.main+xml")
            zout.writestr(item, data)
    dst.write_bytes(buf.getvalue())


def _pptx_template(path: Path) -> None:
    prs = Presentation()  # 4:3 (10 x 7.5 인치)
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "기존 표지 - 지워져야 함"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "기존 내용 - 지워져야 함"
    prs.save(str(path))


def _upload_asset(client, user, partner_id: int, kind: str, path: Path):
    with path.open("rb") as f:
        return client.post(f"/api/partners/{partner_id}/assets/{kind}", files={"file": (path.name, f)}, headers=user)


def _approved_guide(client, user, vendor_id: int, tmp_path: Path) -> int:
    src = tmp_path / "vendor.docx"
    _make_vendor_docx(src)
    with src.open("rb") as f:
        doc_id = client.post("/api/documents", data={"vendor_id": vendor_id}, files={"file": ("vendor.docx", f)}, headers=user).json()["id"]
    run_pending_jobs()
    guide_id = client.post(f"/api/documents/{doc_id}/guides", json={"target_language": "ko"}, headers=user).json()["id"]
    run_pending_jobs()
    client.post(f"/api/guides/{guide_id}/approve-all", headers=user)
    return guide_id


def _export(client, user, guide_id: int, out_dir: Path, **body) -> Path:
    r = client.post(f"/api/guides/{guide_id}/exports", json=body, headers=user)
    assert r.status_code == 201, r.text
    run_pending_jobs()
    export = next(e for e in client.get(f"/api/guides/{guide_id}/exports", headers=user).json() if e["id"] == r.json()["id"])
    assert export["status"] == "done", export
    out = out_dir / f"export_{export['id']}.{body['format']}"
    out.write_bytes(client.get(f"/api/exports/{export['id']}/download", headers=user).content)
    return out


def test_partner_templates(client, tmp_path):
    admin = login(client, *ADMIN)
    partner, vendor = _partner_with_vendor(client, admin, "양식파트너", "TemplateVendor")
    user = _signup_and_approve(client, admin, "tmpl@tmpl.example.com", partner["id"], role="partner_admin")
    guide_id = _approved_guide(client, user, vendor["id"], tmp_path)

    # 양식이 없으면 쓸 수 없다고 알려 준다.
    r = client.post(f"/api/guides/{guide_id}/exports", json={"format": "docx", "use_partner_template": True}, headers=user)
    assert r.status_code == 400 and "Word 양식이 없습니다" in r.json()["detail"]

    logo = tmp_path / "logo.png"
    Image.new("RGB", (120, 40), "darkgreen").save(logo)
    docx_tpl, dotx_tpl, pptx_tpl = tmp_path / "t.docx", tmp_path / "t.dotx", tmp_path / "t.pptx"
    _docx_template(docx_tpl, logo)
    _to_dotx(docx_tpl, dotx_tpl)
    _pptx_template(pptx_tpl)
    assert _upload_asset(client, user, partner["id"], "docx_template", dotx_tpl).status_code == 200
    assert _upload_asset(client, user, partner["id"], "pptx_template", pptx_tpl).status_code == 200

    # Word: 기존 본문·표는 지워지고, 머리글(로고 포함)은 남고, 요약이 들어간다.
    out = Document(str(_export(client, user, guide_id, tmp_path, format="docx", use_partner_template=True)))
    body_text = "\n".join(p.text for p in out.paragraphs) + "".join(c.text for t in out.tables for r in t.rows for c in r.cells)
    assert "지워져야 함" not in body_text
    assert "[모의 요약" in body_text and "문서 버전: 1.0" in body_text
    header = out.sections[0].header
    assert "파트너 머리글" in header.paragraphs[0].text
    assert any("image" in rel.reltype for rel in header.part.rels.values())

    # PPT: 기존 슬라이드는 지워지고, 양식의 슬라이드 크기(4:3)와 레이아웃을 쓴다.
    prs = Presentation(str(_export(client, user, guide_id, tmp_path, format="pptx", use_partner_template=True)))
    texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
    assert not any("지워져야 함" in t for t in texts)
    assert (prs.slide_width, prs.slide_height) == (Inches(10), Inches(7.5))
    assert prs.slides[0].slide_layout.name == "Title Slide"
    assert prs.slides[1].slide_layout.name == "Title and Content"
    assert len(prs.slides) >= 3

    # 관리자가 올린 문서(파트너사 없음)에도 파트너사를 골라 그 양식을 쓸 수 있다.
    admin_guide = _approved_guide(client, admin, vendor["id"], tmp_path)
    out = Document(str(_export(client, admin, admin_guide, tmp_path, format="docx", use_partner_template=True, partner_id=partner["id"])))
    assert "파트너 머리글" in out.sections[0].header.paragraphs[0].text

    # 다른 파트너사의 양식은 쓸 수 없다.
    other_partner, _ = _partner_with_vendor(client, admin, "양식다른", "TemplateOther")
    r = client.post(
        f"/api/guides/{guide_id}/exports",
        json={"format": "docx", "use_partner_template": True, "partner_id": other_partner["id"]},
        headers=user,
    )
    assert r.status_code == 403
