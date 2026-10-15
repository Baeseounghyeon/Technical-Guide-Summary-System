"""가입 승인 → 업로드 → 목차 확인 → 요약 → 검토 → Word/PPT 생성, 도메인 격리까지 한 번에 검증한다."""

from pathlib import Path

from docx import Document
from PIL import Image

from app.worker import run_pending_jobs

from .conftest import ADMIN, login


def _make_vendor_docx(path: Path) -> None:
    d = Document()
    d.add_heading("Acme Router Install Guide", 0)
    d.add_paragraph("Copyright 2026 Acme. All rights reserved.")
    d.add_heading("Overview", 1)
    d.add_paragraph("The Acme router connects branch offices.")
    d.add_heading("Installation", 1)
    d.add_paragraph("Prerequisites: 8 GB RAM and firmware 4.2.")
    d.add_heading("Install steps", 2)
    d.add_paragraph("1. Run `setup.exe /quiet`")
    d.add_paragraph("2. Open port 8443")
    d.add_paragraph("Warning: the device reboots after installation.")
    d.save(str(path))


def _signup_and_approve(client, admin, email: str, partner_id: int, role: str = "engineer") -> dict:
    password = "engineer-pass-1"
    r = client.post(
        "/api/auth/signup",
        json={"email": email, "password": password, "name": "엔지니어", "company_name": "회사", "department": "기술지원"},
    )
    assert r.status_code == 201, r.text
    user_id = r.json()["id"]
    assert client.post("/api/auth/login", json={"email": email, "password": password}).status_code == 403

    r = client.post(f"/api/admin/users/{user_id}/approve", json={"partner_id": partner_id, "role": role}, headers=admin)
    assert r.status_code == 200, r.text
    return login(client, email, password)


def _partner_with_vendor(client, admin, partner_name: str, vendor_name: str) -> tuple[dict, dict]:
    partner = client.post("/api/admin/partners", json={"name": partner_name}, headers=admin).json()
    vendor = client.post("/api/admin/vendors", json={"name": vendor_name}, headers=admin).json()
    r = client.put(f"/api/admin/partners/{partner['id']}/vendors", json={"vendor_ids": [vendor["id"]]}, headers=admin)
    assert r.status_code == 200, r.text
    return partner, vendor


def test_full_flow(client, tmp_path):
    admin = login(client, *ADMIN)
    partner, vendor = _partner_with_vendor(client, admin, "파트너A", "Acme")
    user = _signup_and_approve(client, admin, "eng-a@partner-a.example.com", partner["id"], role="partner_admin")

    me = client.get("/api/auth/me", headers=user).json()
    assert [v["name"] for v in me["vendors"]] == ["Acme"]

    # 파트너 로고 등록
    logo = tmp_path / "logo.png"
    Image.new("RGB", (120, 40), "navy").save(logo)
    with logo.open("rb") as f:
        r = client.post(f"/api/partners/{partner['id']}/assets/logo", files={"file": ("logo.png", f, "image/png")}, headers=user)
    assert r.status_code == 200, r.text

    # 1) 업로드 → 파싱
    src = tmp_path / "acme.docx"
    _make_vendor_docx(src)
    with src.open("rb") as f:
        r = client.post("/api/documents", data={"vendor_id": vendor["id"]}, files={"file": ("acme.docx", f)}, headers=user)
    assert r.status_code == 201, r.text
    doc_id = r.json()["id"]
    assert run_pending_jobs() == 1

    doc = client.get(f"/api/documents/{doc_id}", headers=user).json()
    assert doc["status"] == "parsed"
    sections = {s["title"]: s for s in doc["sections"]}
    assert {"Overview", "Installation", "Install steps"} <= sections.keys()
    assert sections["(표지·서문)"]["included"] is False  # 표지·저작권은 기본 제외
    assert sections["Install steps"]["parent_id"] == sections["Installation"]["id"]

    # 2) 요약 + 번역
    r = client.post(f"/api/documents/{doc_id}/guides", json={"target_language": "ko"}, headers=user)
    assert r.status_code == 201, r.text
    guide_id = r.json()["id"]
    run_pending_jobs()

    guide = client.get(f"/api/guides/{guide_id}", headers=user).json()
    assert guide["status"] == "reviewing"
    assert guide["progress"]["generated"] == guide["progress"]["total"] == 3
    assert guide["overview"]
    summaries = [it["summary"] for it in guide["items"] if it["summary"]]
    steps = next(it["summary"] for it in guide["items"] if it["title"] == "Install steps")
    assert steps["key_steps"] and steps["warnings"]

    # 3) 승인 전에는 최종 문서를 만들 수 없다.
    r = client.post(f"/api/guides/{guide_id}/exports", json={"format": "docx"}, headers=user)
    assert r.status_code == 409

    # 4) 피드백: 코멘트로 재생성, 직접 수정
    first = summaries[0]["id"]
    base = f"/api/guides/{guide_id}/summaries/{first}"
    assert client.post(f"{base}/regenerate", json={"comment": "사전 조건을 더 자세히"}, headers=user).status_code == 200
    run_pending_jobs()
    r = client.put(base, json={"content": "직접 수정한 요약"}, headers=user)
    assert r.json()["version"] == 3
    revisions = client.get(f"{base}/revisions", headers=user).json()
    assert [rev["origin"] for rev in revisions] == ["generated", "regenerated", "edited"]
    assert revisions[1]["comment"] == "사전 조건을 더 자세히"

    for s in summaries:
        r = client.post(f"/api/guides/{guide_id}/summaries/{s['id']}/approve", headers=user)
        assert r.status_code == 200, r.text
    assert client.get(f"/api/guides/{guide_id}", headers=user).json()["status"] == "approved"

    # 5) Word / PPT 생성 (로고 + 이력 관리)
    revision = {"company": "파트너A", "department": "기술지원", "engineer": "홍길동",
                "phone": "010-0000-0000", "email": "eng-a@partner-a.example.com", "change_note": "최초 작성"}
    for fmt in ("docx", "pptx"):
        r = client.post(
            f"/api/guides/{guide_id}/exports",
            json={"format": fmt, "doc_version": "1.0", "include_logo": True, "revision": revision},
            headers=user,
        )
        assert r.status_code == 201, r.text
    run_pending_jobs()

    exports = client.get(f"/api/guides/{guide_id}/exports", headers=user).json()
    assert [(e["format"], e["status"]) for e in exports] == [("docx", "done"), ("pptx", "done")], exports
    for e in exports:
        r = client.get(f"/api/exports/{e['id']}/download", headers=user)
        assert r.status_code == 200 and r.content[:2] == b"PK"

    word = Document(str(Path(_export_path(exports[0]["id"]))))
    text = "\n".join(p.text for p in word.paragraphs)
    assert "직접 수정한 요약" in text and "주의 사항" in text
    assert word.tables[0].rows[1].cells[4].text == "홍길동"
    assert len(word.tables[0].rows) == 2  # Word·PPT 를 같은 버전으로 만들어도 이력은 한 줄
    assert "언어: 한국어" in text

    # 6) 도메인 격리: 다른 벤더 도메인의 파트너는 볼 수도, 올릴 수도 없다.
    partner_b, _ = _partner_with_vendor(client, admin, "파트너B", "Globex")
    other = _signup_and_approve(client, admin, "eng-b@partner-b.example.com", partner_b["id"])
    assert client.get(f"/api/documents/{doc_id}", headers=other).status_code == 404
    assert client.get(f"/api/guides/{guide_id}", headers=other).status_code == 404
    assert client.get("/api/documents", headers=other).json() == []
    assert client.get(f"/api/exports/{exports[0]['id']}/download", headers=other).status_code == 404
    with src.open("rb") as f:
        r = client.post("/api/documents", data={"vendor_id": vendor["id"]}, files={"file": ("acme.docx", f)}, headers=other)
    assert r.status_code == 403


def _export_path(export_id: int) -> str:
    from app.db import SessionLocal
    from app.models import Export

    with SessionLocal() as db:
        return db.get(Export, export_id).file_path
