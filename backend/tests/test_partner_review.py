"""2단계: 파트너 가이드 업로드 → 목차 매칭 → 오류 탐지 → 수락/거절/되돌리기 → 수정본·보고서, 도메인 격리."""

from pathlib import Path

from docx import Document

from app.worker import run_pending_jobs

from .conftest import ADMIN, login
from .test_flow import _partner_with_vendor, _signup_and_approve


def _vendor_docx(path: Path) -> None:
    d = Document()
    d.add_heading("Installation", 1)
    d.add_paragraph("Prerequisites: firmware 4.2 or later.")
    d.add_paragraph("Reserve TCP port 8443 for management.")
    d.add_paragraph("Run `aer-cli setup --mode zero-touch`.")
    d.add_heading("Upgrade", 1)
    d.add_paragraph("Back up the configuration with `aer-cli backup`.")
    d.save(str(path))


def _partner_docx(path: Path) -> None:
    d = Document()
    d.add_heading("설치 가이드", 1)
    d.add_paragraph("사전 조건: firmware 4.1 이상이 필요합니다.")
    d.add_paragraph("관리용 port 8080 을 개방합니다.")
    d.add_paragraph("`aer-cli setup --mode manual` 을 실행합니다.")
    d.add_heading("업그레이드", 1)
    d.add_paragraph("`aer-cli backup` 으로 설정을 백업합니다.")
    d.add_heading("회사 연락처", 1)
    d.add_paragraph("문의: 기술지원팀")
    d.save(str(path))


def _approved_guide(client, user, vendor_id: int, tmp_path: Path) -> int:
    src = tmp_path / "vendor.docx"
    _vendor_docx(src)
    with src.open("rb") as f:
        doc_id = client.post("/api/documents", data={"vendor_id": vendor_id}, files={"file": ("vendor.docx", f)}, headers=user).json()["id"]
    run_pending_jobs()
    guide_id = client.post(f"/api/documents/{doc_id}/guides", json={"target_language": "ko"}, headers=user).json()["id"]
    run_pending_jobs()
    return guide_id


def _approve_all(client, user, guide_id: int) -> None:
    guide = client.get(f"/api/guides/{guide_id}", headers=user).json()
    for item in guide["items"]:
        if item["summary"]:
            client.post(f"/api/guides/{guide_id}/summaries/{item['summary']['id']}/approve", headers=user)


def _upload_partner(client, user, guide_id: int, path: Path):
    with path.open("rb") as f:
        return client.post(f"/api/guides/{guide_id}/partner-guides", files={"file": ("partner.docx", f)}, headers=user)


def test_partner_review_flow(client, tmp_path):
    admin = login(client, *ADMIN)
    partner, vendor = _partner_with_vendor(client, admin, "리뷰파트너", "ReviewVendor")
    user = _signup_and_approve(client, admin, "reviewer@review.example.com", partner["id"])
    guide_id = _approved_guide(client, user, vendor["id"], tmp_path)
    partner_file = tmp_path / "partner.docx"
    _partner_docx(partner_file)

    # 승인 전에는 비교 기준으로 쓸 수 없다.
    assert _upload_partner(client, user, guide_id, partner_file).status_code == 409
    _approve_all(client, user, guide_id)

    # 1) 업로드 → 파싱 + 목차 자동 매칭 (한국어 제목 ↔ 영어 원본)
    r = _upload_partner(client, user, guide_id, partner_file)
    assert r.status_code == 201, r.text
    pg_id = r.json()["id"]
    run_pending_jobs()
    detail = client.get(f"/api/partner-guides/{pg_id}", headers=user).json()
    assert detail["status"] == "matched", detail
    matched = {s["title"]: s["matched_title"] for s in detail["sections"]}
    assert matched == {"설치 가이드": "Installation", "업그레이드": "Upgrade", "회사 연락처": None}

    # 수동 매칭 수정: 다른 문서의 목차는 고를 수 없다.
    contact = next(s for s in detail["sections"] if s["title"] == "회사 연락처")
    assert client.patch(f"/api/partner-sections/{contact['id']}", json={"matched_section_id": 999999}, headers=user).status_code == 400

    # 2) 오류 탐지 (근거는 벤더 원본)
    assert client.post(f"/api/partner-guides/{pg_id}/compare", headers=user).status_code == 200
    run_pending_jobs()
    detail = client.get(f"/api/partner-guides/{pg_id}", headers=user).json()
    assert detail["status"] == "reviewing"
    findings = {f["category"] + ":" + f["before_text"]: f for f in detail["findings"]}
    firmware = next(f for f in detail["findings"] if "4.1" in f["before_text"])
    port = next(f for f in detail["findings"] if "8080" in f["before_text"])
    command = next(f for f in detail["findings"] if f["category"] == "wrong_command")
    assert len(findings) == 3
    assert "4.2" in firmware["after_text"] and "8443" in port["after_text"]
    assert "zero-touch" in command["after_text"] and all(f["locatable"] for f in detail["findings"])
    assert firmware["vendor_section_title"] == "Installation"

    # 3) 결정: 수락 / 수정해서 수락 / 거절(사유 필수) / 되돌리기
    decide = lambda fid, **body: client.post(f"/api/findings/{fid}/decision", json=body, headers=user)  # noqa: E731
    assert decide(firmware["id"], action="accept").status_code == 200
    assert decide(port["id"], action="accept", final_text="관리용 TCP port 8443 을 개방합니다.", note="TCP 명시").status_code == 200
    assert decide(command["id"], action="reject").status_code == 400
    assert decide(command["id"], action="reject", note="파트너 환경은 수동 모드 사용").status_code == 200
    assert decide(command["id"], action="accept").status_code == 409  # 되돌리기 전에는 다시 결정 불가
    assert decide(command["id"], action="reopen").status_code == 200
    assert decide(command["id"], action="reject", note="수동 모드 유지").json()["status"] == "rejected"

    history = client.get(f"/api/partner-guides/{pg_id}/history", headers=user).json()
    assert [h["action"] for h in reversed(history)] == ["accepted", "accepted", "rejected", "reopened", "rejected"]
    assert history[0]["user_name"] == "엔지니어" and history[0]["note"] == "수동 모드 유지"

    # 다시 비교해도 이미 결정한 후보는 중복으로 생기지 않는다.
    client.post(f"/api/partner-guides/{pg_id}/compare", headers=user)
    run_pending_jobs()
    assert len(client.get(f"/api/partner-guides/{pg_id}", headers=user).json()["findings"]) == 3

    # 4) 변경 내용 추적 수정본 + 수정 보고서
    for kind in ("tracked", "report"):
        assert client.post(f"/api/partner-guides/{pg_id}/exports", json={"kind": kind}, headers=user).status_code == 201
    run_pending_jobs()
    exports = client.get(f"/api/partner-guides/{pg_id}/exports", headers=user).json()
    assert [(e["kind"], e["status"]) for e in exports] == [("tracked", "done"), ("report", "done")], exports
    assert exports[0]["applied_count"] == 2 and exports[0]["skipped_count"] == 0

    tracked = client.get(f"/api/correction-exports/{exports[0]['id']}/download", headers=user)
    assert tracked.status_code == 200
    tracked_path = tmp_path / "tracked.docx"
    tracked_path.write_bytes(tracked.content)
    xml = Document(str(tracked_path)).element.body.xml
    assert "w:delText" in xml and "w:ins " in xml
    assert "관리용 TCP port 8443 을 개방합니다." in xml  # 사용자가 고친 문구가 반영됨
    assert "firmware 4.2" in xml and 'w:author="엔지니어"' in xml

    report_path = tmp_path / "report.docx"
    report_path.write_bytes(client.get(f"/api/correction-exports/{exports[1]['id']}/download", headers=user).content)
    table = Document(str(report_path)).tables[0]
    assert len(table.rows) == 4
    assert any("수동 모드 유지" in c.text for c in table.rows[-1].cells + table.rows[1].cells + table.rows[2].cells + table.rows[3].cells)

    # 5) 도메인 격리
    partner_b, _ = _partner_with_vendor(client, admin, "다른파트너", "OtherVendor")
    other = _signup_and_approve(client, admin, "other@other.example.com", partner_b["id"])
    assert client.get(f"/api/partner-guides/{pg_id}", headers=other).status_code == 404
    assert client.post(f"/api/findings/{firmware['id']}/decision", json={"action": "reopen"}, headers=other).status_code == 404
    assert client.get(f"/api/correction-exports/{exports[0]['id']}/download", headers=other).status_code == 404
