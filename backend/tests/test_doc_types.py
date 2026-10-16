"""문서 종류(규격서·기술 가이드·브로슈어): 등록할 때 고르고, 종류·벤더별로 목록을 본다."""

from pathlib import Path

from app.worker import run_pending_jobs

from .conftest import ADMIN, login
from .test_flow import _make_vendor_docx, _partner_with_vendor, _signup_and_approve


def test_documents_by_type_and_vendor(client, tmp_path: Path):
    admin = login(client, *ADMIN)
    partner, vendor = _partner_with_vendor(client, admin, "종류파트너", "TypeVendorA")
    _, vendor_b = _partner_with_vendor(client, admin, "종류파트너B", "TypeVendorB")
    client.put(f"/api/admin/partners/{partner['id']}/vendors", json={"vendor_ids": [vendor["id"], vendor_b["id"]]}, headers=admin)
    user = _signup_and_approve(client, admin, "types@t.example.com", partner["id"])
    src = tmp_path / "v.docx"
    _make_vendor_docx(src)

    def upload(vendor_id: int, doc_type: str | None, title: str) -> dict:
        data = {"vendor_id": vendor_id, "title": title, "description": f"{title} 설명"}
        if doc_type:
            data["doc_type"] = doc_type
        with src.open("rb") as f:
            r = client.post("/api/documents", data=data, files={"file": ("v.docx", f)}, headers=user)
        assert r.status_code == 201, r.text
        return r.json()

    spec = upload(vendor["id"], "spec", "A 규격서")
    assert spec["doc_type"] == "spec" and spec["description"] == "A 규격서 설명" and spec["uploaded_by_name"]
    upload(vendor["id"], None, "A 가이드")  # 종류를 안 고르면 기술 가이드
    upload(vendor_b["id"], "brochure", "B 브로슈어")
    run_pending_jobs()  # 다른 테스트에 대기 작업을 남기지 않는다

    def titles(query: str) -> list[str]:
        return sorted(d["title"] for d in client.get(f"/api/documents?{query}", headers=user).json())

    assert titles("doc_type=spec") == ["A 규격서"]
    assert titles("doc_type=guide") == ["A 가이드"]
    assert titles(f"doc_type=brochure&vendor_id={vendor['id']}") == []
    assert titles(f"doc_type=brochure&vendor_id={vendor_b['id']}") == ["B 브로슈어"]
    assert client.post(
        "/api/documents", data={"vendor_id": vendor["id"], "doc_type": "unknown", "url": "https://example.com"}, headers=user
    ).status_code == 422


def test_guide_kinds(client, tmp_path: Path):
    """기술 가이드 문서는 기술 가이드 / 트러블 슈팅 가이드로 나뉜다 (가이드 화면의 종류 선택)."""
    admin = login(client, *ADMIN)
    partner, vendor = _partner_with_vendor(client, admin, "가이드종류파트너", "KindVendor")
    user = _signup_and_approve(client, admin, "kinds@k.example.com", partner["id"])
    src = tmp_path / "v.docx"
    _make_vendor_docx(src)
    for title, kind in [("설치 가이드", "technical"), ("장애 대응", "troubleshooting"), ("기본값", None)]:
        data = {"vendor_id": vendor["id"], "doc_type": "guide", "title": title}
        if kind:
            data["guide_kind"] = kind
        with src.open("rb") as f:
            assert client.post("/api/documents", data=data, files={"file": ("v.docx", f)}, headers=user).status_code == 201
    with src.open("rb") as f:
        spec = client.post(
            "/api/documents", data={"vendor_id": vendor["id"], "doc_type": "spec", "guide_kind": "troubleshooting", "title": "규격"},
            files={"file": ("v.docx", f)}, headers=user,
        ).json()
    assert spec["guide_kind"] is None  # 가이드가 아닌 문서에는 종류가 없다
    run_pending_jobs()

    def titles(query: str) -> list[str]:
        return sorted(d["title"] for d in client.get(f"/api/documents?doc_type=guide&vendor_id={vendor['id']}{query}", headers=user).json())

    assert titles("") == ["기본값", "설치 가이드", "장애 대응"]
    assert titles("&guide_kind=technical") == ["기본값", "설치 가이드"]
    assert titles("&guide_kind=troubleshooting") == ["장애 대응"]


def test_download_formats_delete_and_filters(client, tmp_path: Path):
    """원본·Word·PPT·Excel·PDF 내려받기, 파트너·작성자 필터, 관리자 삭제."""
    import io

    import pymupdf
    from docx import Document
    from openpyxl import load_workbook

    admin = login(client, *ADMIN)
    partner, vendor = _partner_with_vendor(client, admin, "다운로드파트너", "DownloadVendor")
    user = _signup_and_approve(client, admin, "dl@d.example.com", partner["id"])
    src = tmp_path / "v.docx"
    _make_vendor_docx(src)
    with src.open("rb") as f:
        doc = client.post("/api/documents", data={"vendor_id": vendor["id"], "title": "다운로드 가이드"}, files={"file": ("v.docx", f)}, headers=user).json()
    run_pending_jobs()
    base = f"/api/documents/{doc['id']}/download"

    assert client.get(f"{base}?format=original", headers=user).content == src.read_bytes()
    word = Document(io.BytesIO(client.get(f"{base}?format=docx", headers=user).content))
    assert "Install steps" in "\n".join(p.text for p in word.paragraphs)
    assert client.get(f"{base}?format=pptx", headers=user).content[:2] == b"PK"
    sheet = load_workbook(io.BytesIO(client.get(f"{base}?format=xlsx", headers=user).content)).active
    assert sheet["C1"].value == "목차" and any(r[2] == "Install steps" for r in sheet.iter_rows(values_only=True))
    r = client.get(f"{base}?format=pdf", headers=user)
    assert "attachment" in r.headers["content-disposition"]
    assert "Install steps" in "".join(p.get_text() for p in pymupdf.open("pdf", r.content))
    assert client.get(f"{base}?format=exe", headers=user).status_code == 422

    # 파트너·작성자 필터
    me_id = client.get("/api/auth/me", headers=user).json()["id"]
    listed = client.get(f"/api/documents?doc_type=guide&partner_id={partner['id']}&uploaded_by={me_id}", headers=admin).json()
    assert [d["id"] for d in listed] == [doc["id"]] and listed[0]["partner_name"] == "다운로드파트너"
    assert client.get(f"/api/documents?uploaded_by={me_id + 999}", headers=admin).json() == []

    # 삭제는 플랫폼 관리자만
    assert client.delete(f"/api/documents/{doc['id']}", headers=user).status_code == 403
    assert client.delete(f"/api/documents/{doc['id']}", headers=admin).status_code == 204
    assert client.get(f"/api/documents/{doc['id']}", headers=admin).status_code == 404
