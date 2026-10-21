"""견적서 생성·생성한 견적서. 업무 구분: 견적서 생성은 영업만(관리자 포함), 엔지니어는 불가.
(규격서 생성 기능은 삭제됨 — 그 API 가 없어졌는지도 확인한다.)"""

import io
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import load_workbook

from app.worker import run_pending_jobs

from .conftest import ADMIN, login
from .test_flow import _partner_with_vendor, _signup_and_approve


def test_spec_generation_removed(client):
    """규격서 생성·생성한 규격서는 삭제: 그 API 는 없고, 목록도 견적서만 받는다."""
    admin = login(client, *ADMIN)
    assert client.get("/api/forms/spec/fields", headers=admin).status_code == 404
    assert client.post("/api/generated", json={"doc_type": "spec", "title": "x", "fields": {}}, headers=admin).status_code == 405
    assert client.get("/api/generated?doc_type=spec", headers=admin).status_code == 422


def test_my_guides_lists_created_guides(client):
    admin = login(client, *ADMIN)
    assert isinstance(client.get("/api/my-guides", headers=admin).json(), list)


def _xlsx(path: Path, rows: list[tuple]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "규격"
    for r in rows:
        ws.append(list(r))
    wb.save(str(path))


def _sales(client, admin, email: str, partner_id: int) -> dict:
    headers = _signup_and_approve(client, admin, email, partner_id)
    user_id = client.get("/api/auth/me", headers=headers).json()["id"]
    client.post(f"/api/admin/users/{user_id}/approve", json={"partner_id": partner_id, "track": "sales"}, headers=admin)
    return headers


def test_price_list_purchase_types(client, tmp_path: Path):
    """제품 가격표·매입 금액: 문서 등록의 종류(영업만, PDF·Word·Excel). 규격서도 Excel 을 받는다."""
    admin = login(client, *ADMIN)
    partner, vendor = _partner_with_vendor(client, admin, "가격표파트너", "PriceVendor")
    sales = _sales(client, admin, "price-sales@p.example.com", partner["id"])
    engineer = _signup_and_approve(client, admin, "price-eng@p.example.com", partner["id"])
    assert {"price_list", "purchase"} <= set(client.get("/api/auth/me", headers=sales).json()["permissions"])
    assert not {"price_list", "purchase"} & set(client.get("/api/auth/me", headers=engineer).json()["permissions"])

    def upload(headers, name: str, data: dict):
        with (tmp_path / name).open("rb") as f:
            return client.post("/api/documents", data={"vendor_id": vendor["id"], **data}, files={"file": (name, f)}, headers=headers)

    _xlsx(tmp_path / "spec.xlsx", [("항목 A", "Luna HSM"), ("항목 B", "A790"), ("항목 C", "2대"), ("항목 D", "본사"), ("항목 E", "FIPS 140-3")])
    spec = upload(sales, "spec.xlsx", {"doc_type": "spec", "title": "엑셀 규격서"}).json()
    assert spec["source_type"] == "xlsx"
    assert upload(sales, "spec.xlsx", {"doc_type": "brochure"}).status_code == 400  # Excel 은 규격서·가격표·매입만

    _xlsx(tmp_path / "price.xlsx", [("제품", "단가"), ("Luna HSM", 30000000)])
    assert upload(engineer, "price.xlsx", {"doc_type": "price_list"}).status_code == 403
    price = upload(sales, "price.xlsx", {"doc_type": "price_list", "title": "가격표"}).json()
    purchase = upload(sales, "price.xlsx", {"doc_type": "purchase", "title": "매입가"}).json()
    assert client.post(
        "/api/documents", data={"vendor_id": vendor["id"], "doc_type": "purchase", "url": "https://example.com"}, headers=sales
    ).status_code == 400  # 파일만
    run_pending_jobs()

    assert [d["title"] for d in client.get("/api/documents?doc_type=price_list", headers=sales).json()] == ["가격표"]



def _template_docx(path: Path) -> None:
    d = Document()
    d.add_paragraph("견적서: {{제목}}")
    d.add_paragraph("작성자 {{작성자}} / {{소속}} / {{전화번호}}")
    d.add_paragraph("{{제품가격표}}")
    d.add_paragraph("매입 내역")
    d.add_paragraph("{{매입금액}}")
    d.save(str(path))


def _thales_price_list_pdf(path: Path) -> None:
    """Thales 가격표와 같은 배치: SKU | 제품명 | 설명 | $ 가격 이 한 줄."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((112, 60), "Thales CSP HSM Price List", fontsize=8)
    for y, (part, name, desc, price) in enumerate([
        ("908-000364-003-000", "Luna Network HSM S750", "Luna Network HSM S750 (PED,Enterprise perf,16MB,5 Partitions)", "$      55,176"),
        ("020-160001-002-000", "HSM", "HSM,Enhanced Maintenance Service,1-2 Year", "21.0%"),
    ]):
        for x, text in ((112, part), (170, name), (286, desc), (491, price)):
            page.insert_text((x, 100 + y * 12), text, fontsize=5)
    doc.save(str(path))


def _thales_quote_pdf(path: Path, number: str, lines: list[tuple], total: str) -> None:
    """Thales 매입 견적과 같은 머리 정보와 라인 표."""
    head = ["Line Number", "Maintenance Part Number / Description", "Covered Product Part Number / Description",
            "Serial Number / External Reference", "Service Coverage", "Quantity", "Maintenance Unit Price", "Net Line Total",
            "SFDC / Oracle Contract Number"]
    cell = "border:1px solid black;padding:2px;font-size:6pt"
    rows = "".join("<tr>" + "".join(f'<td style="{cell}">{v}</td>' for v in r) + "</tr>" for r in lines)
    html = (
        f"<p>Quote Number: {number}</p><p>Date Quoted:</p><p>2026-09-21</p>"
        "<p>Primary Partner Information</p><p>Partner Co., Ltd.</p>"
        "<p>End User Account Name:</p><p>EndUser Contact:</p><p>Example Bank Inc.</p>"
        f'<table style="border-collapse:collapse"><tr>{"".join(f"<td style=\'{cell}\'>{h}</td>" for h in head)}</tr>{rows}</table>'
        f"<p>Quote Total: USD {total}</p>"
    )
    story = pymupdf.Story(html=html)
    writer = pymupdf.DocumentWriter(str(path))
    page = pymupdf.paper_rect("a4-l")
    more = True
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(page + (36, 36, -36, -36))
        story.draw(device)
        writer.end_page()
    writer.close()


def test_quote_builder_steps(client, tmp_path: Path):
    """견적서_매핑정리.md 기준: 매입 견적 PDF 한 건 = 견적서 한 행, N=Net Line Total(여러 라인이면 합), G=N×O×2.3, 목표 금액으로 할인율 역산."""
    admin = login(client, *ADMIN)
    partner, vendor = _partner_with_vendor(client, admin, "견적생성파트너", "QuoteBuildVendor")
    sales = _sales(client, admin, "qb-sales@q.example.com", partner["id"])
    engineer = _signup_and_approve(client, admin, "qb-eng@q.example.com", partner["id"])

    # 1. 견적서 양식 (Word·Excel 만, 영업만)
    _template_docx(tmp_path / "tpl.docx")
    with (tmp_path / "tpl.docx").open("rb") as f:
        tpl = client.post("/api/quote-templates", data={"name": "기본 견적 양식"}, files={"file": ("tpl.docx", f)}, headers=sales).json()
    with (tmp_path / "tpl.docx").open("rb") as f:
        assert client.post("/api/quote-templates", data={"name": "x"}, files={"file": ("tpl.pdf", f)}, headers=sales).status_code == 400
    with (tmp_path / "tpl.docx").open("rb") as f:
        assert client.post("/api/quote-templates", data={"name": "x"}, files={"file": ("tpl.docx", f)}, headers=engineer).status_code == 403
    assert client.put(f"/api/quote-templates/{tpl['id']}", data={"name": "견적 양식 v2"}, headers=sales).json()["name"] == "견적 양식 v2"

    def upload(path: Path, doc_type: str, title: str) -> dict:
        with path.open("rb") as f:
            r = client.post("/api/documents", data={"vendor_id": vendor["id"], "doc_type": doc_type, "title": title}, files={"file": (path.name, f)}, headers=sales)
        assert r.status_code == 201, r.text
        return r.json()

    # 2·3. 가격표 (참고용): 제품 번호·제품명·설명·리스트 가격만
    _thales_price_list_pdf(tmp_path / "price.pdf")
    price = upload(tmp_path / "price.pdf", "price_list", "HSM 가격표")
    # 4·5. 매입 견적 2건: 한 건은 라인 1개, 다른 건은 라인 2개 (대표 라인 + 다른 라인)
    _thales_quote_pdf(tmp_path / "q1.pdf", "Q-1001", [
        ("1", "020-160001-002-000/HSM, ENHANCED MAINTENANCE SERVICE, 1-2 YEAR",
         "908-000364-003-000 / LUNA NETWORK HSM S750 (PED, ENTERPRISE PERF, 16MB, GRK-16)", "AB100001, AB100002 (ABCD1-100001)",
         "12 months", "2.00", "USD 7,337.84", "USD 14,675.68", ""),
    ], "14,675.68")
    _thales_quote_pdf(tmp_path / "q2.pdf", "Q-1002", [
        ("1", "020-160001-002-000/HSM, ENHANCED MAINTENANCE SERVICE, 1-2 YEAR", "908-000393-001-000 / KOREAN ALGORITHMS, LUNA NETWORK HSM 7",
         "ABCD1-100002", "12 months StartDate = 01-January-2027 EndDate = 31-December-2027", "6.00", "USD 405.87", "USD 2,435.22", ""),
        ("7", "020-160001-002-000/HSM, ENHANCED MAINTENANCE SERVICE, 1-2 YEAR", "908-000364-003-000 / LUNA NETWORK HSM S750 (PED, ENTERPRISE PERF)",
         "ABCD1-100002 S/N: AB200001, AB200002", "12 months StartDate = 01-January-2027 EndDate = 31-December-2027", "6.00", "USD 3,039.34",
         "USD 18,236.04", ""),
    ], "20,671.26")
    q1 = upload(tmp_path / "q1.pdf", "purchase", "Q-1001")
    q2 = upload(tmp_path / "q2.pdf", "purchase", "Q-1002")
    run_pending_jobs()

    items = client.get(f"/api/documents/{price['id']}/items", headers=sales).json()
    assert items["columns"] == ["제품 번호", "제품명", "설명", "리스트 가격(USD)"]
    assert items["rows"][0][:2] == ["908-000364-003-000", "Luna Network HSM S750"] and items["rows"][0][3] == "55176"
    assert items["rows"][1][3] == "21.0%"
    p1 = client.get(f"/api/documents/{q1['id']}/items", headers=sales).json()
    assert p1["meta"]["quote_number"] == "Q-1001" and p1["meta"]["partner"] == "Partner Co., Ltd."
    assert p1["meta"]["end_user"] == "Example Bank Inc." and p1["meta"]["quote_total"] == "14675.68"
    row = dict(zip(p1["columns"], p1["rows"][0]))
    assert row["유지보수 제품 번호"] == "020-160001-002-000" and row["시리얼 번호"] == "AB100001 / AB100002"
    assert row["계약 번호"] == "ABCD1-100001" and row["수량"] == "2" and row["합계(USD)"] == "14675.68"
    p2 = client.get(f"/api/documents/{q2['id']}/items", headers=sales).json()
    assert [r[0] for r in p2["rows"]] == ["1", "7"]
    assert dict(zip(p2["columns"], p2["rows"][1]))["시작일"] == "2027.01.01"

    # 6·7. 미리보기 → 생성 (환율 1360, 배수 2.3, 목표 공급가 20,000,000)
    body = {
        "title": "[예시은행] Luna Network HSM S750 8식 외 Renewal License 공급 건",
        "recipient": "파트너사",
        "date": "2026-09-28",
        "template_id": None,
        "price_list_id": price["id"],
        "price": {"columns": items["columns"], "rows": items["rows"][:1]},
        "purchases": [
            {"document_id": q1["id"], "meta": p1["meta"], "rows": p1["rows"], "included": [True], "main": 0},
            {"document_id": q2["id"], "meta": p2["meta"], "rows": p2["rows"], "included": [True, True], "main": 1},
        ],
        "author": {"name": "영업 담당", "department": "솔루션사업본부", "phone": "010-0000-0000", "email": "s@unipoint.co.kr",
                   "company": "예시 주식회사", "address": "[00000] 예시시 예시로 1", "contact": "Tel : 02-000-0000"},
        "settings": {"exchange_rate": 1360, "multiplier": 2.3, "target_total": 20000000},
    }
    preview = client.post("/api/quote-builder/preview", json=body, headers=sales)
    assert preview.status_code == 200, preview.text
    preview = preview.json()
    assert preview["formats"] == ["xlsx", "pdf"]
    lines, totals = preview["lines"], preview["totals"]
    assert [(l["quantity"], l["cost_usd"]) for l in lines] == [(2, 14675.68), (6, 20671.26)]  # 2번 행은 두 라인의 합
    assert lines[1]["description"].startswith("LUNA NETWORK HSM S750") and "AB200001 / AB200002" in lines[1]["description"]
    assert lines[0]["product_name"] == "Luna Network HSM" and lines[0]["note"] == "Enhanced MA\n(ABCD1-100001)"
    assert round(totals["supply"]) == 20000000 and round(totals["vat_included"]) == 22000000
    assert abs(lines[0]["amount"] - 14675.68 * 1360 * 2.3) < 0.01
    assert "파트너사 귀중" in preview["html"] and "이천이백만" in preview["html"]
    assert client.post("/api/quote-builder", json={**body, "recipient": ""}, headers=sales).status_code == 400
    assert client.post("/api/quote-builder", json=body, headers=engineer).status_code == 403

    q = client.post("/api/quote-builder", json=body, headers=sales)
    assert q.status_code == 201, q.text
    q = q.json()
    assert q["formats"] == ["xlsx", "pdf"] and q["version"] == 1
    xlsx = client.get(f"/api/generated/{q['id']}/download?format=xlsx", headers=sales)
    assert xlsx.status_code == 200
    ws = load_workbook(io.BytesIO(xlsx.content)).worksheets[0]
    assert ws["B7"].value == "파트너사 귀중" and ws["B8"].value.startswith("프로젝트명 : [예시은행]")
    # 회사명·주소·연락처·담당자 줄은 작성자 정보로 (양식에는 비어 있음)
    assert ws["B1"].value == "예시 주식회사" and ws["J1"].value == "[00000] 예시시 예시로 1" and ws["J2"].value == "Tel : 02-000-0000"
    assert ws["B2"].value == "솔루션사업본부/ 영업 담당 / 010-0000-0000 / s@unipoint.co.kr"
    # 양식의 도형(제목 상자·회사 상자)과 직인 그림이 그대로 남는다 (openpyxl 이 버리는 도형은 따로 옮긴다)
    import zipfile

    drawing = zipfile.ZipFile(io.BytesIO(xlsx.content)).read("xl/drawings/drawing1.xml").decode()
    assert "QUOTATION" in drawing and "<xdr:sp" in drawing and len(ws._images) == 2  # 로고·직인
    # Product Description: 제품·사양·시리얼은 보통 검은 글씨, '(기간) 1year MA' 는 굵은 빨간 글씨
    rich = load_workbook(io.BytesIO(xlsx.content), rich_text=True).worksheets[0]["D18"].value
    plain, period = rich[0], rich[1]
    assert plain.font.b is False and period.font.b is True and period.font.color.rgb == "FFFF0000"
    assert period.text.endswith("1year MA") and "Serial Num." in plain.text
    assert ws["N18"].value == 14675.68 and ws["O18"].value == 1360 and ws["G18"].value == "=N18*O18*2.3"
    assert ws["H19"].value == "=G19*(1-I19)" and ws["S19"].value == "=R19/H19" and ws["P9"].value == "=U14" and ws["U14"].value == 20000000
    # 양식의 품목 표(18~21행) 안에 채운다: 2건이면 20·21행은 빈 채로, 합계 위치는 그대로
    assert ws["B20"].value is None and ws["B21"].value is None and ws["D22"].value == "- 이하여백 -"
    assert ws["H23"].value == "=SUM(H18:H22)" and ws["H24"].value == "=H23*1.1" and ws["P10"].value == "=1-P9/(G18+G19)"
    assert "NUMBERSTRING(H24,1)" in ws["B16"].value and ws.print_area.endswith("$J$32")
    assert client.get(f"/api/generated/{q['id']}/download?format=pdf", headers=sales).content.startswith(b"%PDF")
    assert client.get(f"/api/generated/{q['id']}/download?format=docx", headers=sales).status_code == 400

    # 9. 고쳐서 다시 생성 → 다음 버전
    v2 = client.post("/api/quote-builder", json={**body, "settings": {"exchange_rate": 1400}, "based_on_id": q["id"]}, headers=sales).json()
    assert v2["version"] == 2 and v2["root_id"] == q["id"]
    # 삭제: 다른 사람은 못 지우고, 만든 사람이 지우면 모든 버전이 함께 지워진다
    assert client.delete(f"/api/generated/{v2['id']}", headers=engineer).status_code == 403  # 같은 조직이어도 만든 사람이 아니면
    assert client.delete(f"/api/generated/{v2['id']}", headers=sales).status_code == 204
    assert [g["id"] for g in client.get("/api/generated?doc_type=quote", headers=sales).json()] == []
    q = client.post("/api/quote-builder", json=body, headers=sales).json()
    # 직접 올린 Word 양식으로 만들면 Word·PDF
    v3 = client.post("/api/quote-builder", json={**body, "template_id": tpl["id"]}, headers=sales).json()
    assert v3["formats"] == ["docx", "pdf"]
    text = "\n".join(p.text for p in Document(io.BytesIO(client.get(f"/api/generated/{v3['id']}/download?format=docx", headers=sales).content)).paragraphs)
    assert "견적서: [예시은행]" in text and "영업 담당" in text


def test_quote_table_grows_only_when_needed():
    """품목이 양식 표(4행)보다 많을 때만 '- 이하여백 -' 위에 행을 늘린다."""
    from app.services import unipoint_quote as uq

    row = ["1", "020-160001-002-000", "HSM, ENHANCED MAINTENANCE SERVICE", "908-000364-003-000", "LUNA NETWORK HSM S750 (PED)",
           "AB100001", "ABCD1-100001", "2027.01.01", "2027.12.31", "12 months", "2", "100", "200"]
    s = uq.QuoteSettings()
    lines, totals = uq.build_lines([uq.Purchase(rows=[row], included=[True], main=0) for _ in range(6)], s)
    q = uq.QuoteDoc(project="p", recipient="r", date="2026-10-19", author={"name": "n"}, lines=lines, totals=totals, settings=s)
    ws = load_workbook(io.BytesIO(uq.build_xlsx(q)[0])).worksheets[0]
    assert [ws[f"A{r}"].value for r in range(18, 24)] == [1, 2, 3, 4, 5, 6]
    assert ws["D24"].value == "- 이하여백 -" and ws["H25"].value == "=SUM(H18:H24)" and ws["H26"].value == "=H25*1.1"
    assert ws.row_dimensions[23].height == ws.row_dimensions[21].height and ws.print_area.endswith("$J$34")


def test_xls_documents(client, tmp_path: Path):
    """예전 엑셀(.xls)도 규격서·제품 가격표·매입 금액으로 등록하고 항목을 뽑는다."""
    import xlwt

    admin = login(client, *ADMIN)
    partner, vendor = _partner_with_vendor(client, admin, "XLS파트너", "XlsVendor")
    sales = _sales(client, admin, "xls-sales@x.example.com", partner["id"])
    book = xlwt.Workbook()
    sheet = book.add_sheet("가격표")
    for r, values in enumerate([("제품", "단가"), ("Luna HSM", 55176), ("CTM", 12000)]):
        for c, v in enumerate(values):
            sheet.write(r, c, v)
    book.save(str(tmp_path / "price.xls"))
    with (tmp_path / "price.xls").open("rb") as f:
        r = client.post("/api/documents", data={"vendor_id": vendor["id"], "doc_type": "price_list", "title": "xls 가격표"}, files={"file": ("price.xls", f)}, headers=sales)
    assert r.status_code == 201, r.text
    doc = r.json()
    assert doc["source_type"] == "xls"
    run_pending_jobs()
    assert client.get(f"/api/documents/{doc['id']}", headers=sales).json()["status"] == "parsed"
    items = client.get(f"/api/documents/{doc['id']}/items", headers=sales).json()
    assert items == {"columns": ["제품", "단가"], "rows": [["Luna HSM", "55176"], ["CTM", "12000"]]}
    with (tmp_path / "price.xls").open("rb") as f:  # 브로슈어는 Excel 을 받지 않는다
        assert client.post("/api/documents", data={"vendor_id": vendor["id"], "doc_type": "brochure"}, files={"file": ("price.xls", f)}, headers=sales).status_code == 400


def test_quote_line_order():
    """품목 순서: 수량 오름차순 → 수량이 같으면 기간이 긴 것이 뒤 → 기간도 같으면 Enhanced MA(계약 번호)가 큰 것이 앞."""
    from app.services import unipoint_quote as uq

    def purchase(qty: str, start: str, end: str, contract: str) -> uq.Purchase:
        row = ["1", "020-160001-002-000", "HSM, ENHANCED MAINTENANCE SERVICE", "908-000364-003-000", "LUNA NETWORK HSM S750 (PED)",
               "", contract, start, end, "12 months", qty, "1", "100"]
        return uq.Purchase(rows=[row], included=[True], main=0)

    lines, _ = uq.build_lines([
        purchase("6", "2027.01.01", "2027.12.31", "ABCD1-100004"),
        purchase("4", "2027.01.01", "2027.12.31", "ABCD1-100003"),
        purchase("4", "2027.01.01", "2027.12.31", "ABCD1-100002"),
        purchase("4", "2026.01.01", "2027.12.31", "ABCD1-100001"),  # 기간이 더 길다
        purchase("2", "2026.10.01", "2027.09.30", "ABCD1-100005"),
    ], uq.QuoteSettings())
    assert [(l["no"], l["quantity"], l["contract"]) for l in lines] == [
        (1, 2, "ABCD1-100005"), (2, 4, "ABCD1-100003"), (3, 4, "ABCD1-100002"), (4, 4, "ABCD1-100001"), (5, 6, "ABCD1-100004"),
    ]


def test_shift_formula_when_rows_added():
    """품목 행을 늘리면 그 아래 칸을 가리키는 참조는 내리고, 마지막 품목 행까지의 범위는 늘린다 (따옴표 안 글자는 그대로)."""
    from app.services.unipoint_quote import shift_formula

    assert shift_formula("=SUM(M18:M21)", 22, 2) == "=SUM(M18:M23)"
    assert shift_formula("=SUM(H18:H22)", 22, 2) == "=SUM(H18:H24)"
    assert shift_formula("=M22*N22", 22, 2) == "=M24*N24"
    assert shift_formula("=N21*O21*2.3", 22, 2) == "=N21*O21*2.3"
    assert shift_formula('="제안금액 : 일금"&NUMBERSTRING(H24,1)&"원정(부가세 포함)"', 22, 2) == '="제안금액 : 일금"&NUMBERSTRING(H26,1)&"원정(부가세 포함)"'


def test_period_from_quote_expiry():
    """매입 견적 PDF 에 기간 날짜가 없으면(예: '12 months'만) 견적 만료일 다음 날부터 커버리지 개월로 본다."""
    import re

    from app.services.thales_pdf import _period

    assert _period([], re.search(r"(\d+)", "12 months"), "2026-09-30") == ("2026.10.01", "2027.09.30")
    assert _period([("01", "January", "2027"), ("31", "December", "2027")], None, "") == ("2027.01.01", "2027.12.31")


def test_quote_domains(client, tmp_path: Path):
    """견적서 생성 도메인: 사용자에게 부여된 Thales·Ahnlab 만 고르고, 문서도 그 도메인 것만 쓴다."""
    admin = login(client, *ADMIN)
    partner = client.post("/api/admin/partners", json={"name": "도메인견적파트너"}, headers=admin).json()
    def vendor(name: str) -> dict:  # 다른 테스트가 먼저 만들었으면 그대로 쓴다
        r = client.post("/api/admin/vendors", json={"name": name}, headers=admin)
        return r.json() if r.status_code < 300 else next(v for v in client.get("/api/auth/me", headers=admin).json()["vendors"] if v["name"] == name)

    thales, ahnlab, other = (vendor(n) for n in ("Thales", "Ahnlab", "OtherDomain"))
    client.put(f"/api/admin/partners/{partner['id']}/vendors", json={"vendor_ids": [thales["id"], other["id"]]}, headers=admin)
    sales = _sales(client, admin, "domain-sales@d.example.com", partner["id"])
    assert {v["name"] for v in client.get("/api/auth/me", headers=sales).json()["vendors"]} == {"Thales", "OtherDomain"}

    _xlsx(tmp_path / "price.xlsx", [("제품", "단가"), ("Luna HSM", 30000000)])
    with (tmp_path / "price.xlsx").open("rb") as f:
        price = client.post(
            "/api/documents", data={"vendor_id": other["id"], "doc_type": "price_list", "title": "다른 도메인 가격표"},
            files={"file": ("price.xlsx", f)}, headers=sales,
        ).json()

    def preview(vendor_id, **extra):
        return client.post("/api/quote-builder/preview", json={"vendor_id": vendor_id, **extra}, headers=sales)

    assert preview(thales["id"]).status_code == 200
    assert preview(ahnlab["id"]).status_code == 403  # 부여되지 않은 도메인
    assert preview(other["id"]).status_code == 403  # 견적서 도메인이 아님
    assert preview(thales["id"], price_list_id=price["id"]).status_code == 400  # 다른 도메인 문서

    # 도메인을 바꿔 부여하면 Ahnlab 으로 만들 수 있다
    client.put(f"/api/admin/partners/{partner['id']}/vendors", json={"vendor_ids": [thales["id"], ahnlab["id"]]}, headers=admin)
    assert preview(ahnlab["id"]).status_code == 200
